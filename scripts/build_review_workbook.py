"""Build the card scenario review workbook: data/scenarios/cards/card_scenario_review.xlsx."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SCENARIOS_DIR = ROOT / "data" / "scenarios" / "cards"
OUTPUT = SCENARIOS_DIR / "card_scenario_review.xlsx"

HEADERS = [
    "#",
    "Scenario File",
    "Card Under Test",
    "Card Name",
    "Aspect",
    "Description",
    "Reachable",
    "Playable Now",
    "Move Count",
    "Is Terminal",
    "PASS / FAIL",
    "Failure Category",
    "Failure Description",
    "Expected Behavior",
    "Tester",
    "Date Tested",
    "Notes",
]

COL_WIDTHS = {
    "A": 4,    # #
    "B": 36,   # Scenario File
    "C": 28,   # Card Under Test
    "D": 30,   # Card Name
    "E": 12,   # Aspect
    "F": 48,   # Description
    "G": 11,   # Reachable
    "H": 13,   # Playable Now
    "I": 11,   # Move Count
    "J": 11,   # Is Terminal
    "K": 14,   # PASS / FAIL
    "L": 22,   # Failure Category
    "M": 52,   # Failure Description
    "N": 52,   # Expected Behavior
    "O": 10,   # Tester
    "P": 14,   # Date Tested
    "Q": 40,   # Notes
}

PASS_FAIL_OPTIONS = ["PASS", "FAIL", "UNTESTED", "N/A"]
CATEGORY_OPTIONS = [
    "Engine Bug",
    "Card Data Wrong",
    "Scenario Broken",
    "Card Not Implemented",
    "UI/Interface Issue",
    "Other",
]

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=10)
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
UNTESTED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
NA_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _load_scenario(filepath: Path) -> dict | None:
    try:
        with open(filepath, encoding="utf-8") as fh:
            return json.load(fh)
    except Exception:
        return None


def _extract_card_name(scenario: dict) -> str:
    card_id = scenario["metadata"]["card_under_test"]
    catalog = scenario.get("state_payload", {}).get("definition", {}).get("catalog", {})
    for card in catalog.get("cards", []):
        if card.get("card_id") == card_id:
            return card.get("name", card_id)
    return card_id


def _extract_aspect(scenario: dict) -> str:
    card_id = scenario["metadata"]["card_under_test"]
    catalog = scenario.get("state_payload", {}).get("definition", {}).get("catalog", {})
    for card in catalog.get("cards", []):
        if card.get("card_id") == card_id:
            return card.get("aspect", "")
    return ""


def main():
    files = sorted(
        [f for f in SCENARIOS_DIR.iterdir() if f.suffix == ".json"],
        key=lambda f: f.name,
    )

    wb = Workbook()

    # ── Sheet 1: Review ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Review"

    for col_idx, (col_letter, width) in enumerate(COL_WIDTHS.items(), 1):
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 28

    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row_idx, filepath in enumerate(files, 2):
        scenario = _load_scenario(filepath)
        if scenario is None:
            continue
        meta = scenario["metadata"]
        tags = set(meta.get("tags", []))

        row_data = [
            row_idx - 1,
            filepath.name,
            meta.get("card_under_test", ""),
            _extract_card_name(scenario),
            _extract_aspect(scenario),
            meta.get("description", ""),
            "Yes" if "reachable" in tags else "No",
            "Yes" if "playable_now" in tags else "No",
            meta.get("move_count", ""),
            "Yes" if meta.get("is_terminal") else "No",
            "UNTESTED",   # PASS / FAIL default
            "",
            "",
            "",
            "",
            "",
            "",
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = ROW_FONT
            cell.border = THIN_BORDER
            if col_idx in (1, 7, 8, 9, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (11, 12, 15, 16):
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

            if col_idx == 11:
                cell.fill = UNTESTED_FILL

    last_row = len(files) + 1

    dv_pass_fail = DataValidation(
        type="list",
        formula1=f'"{",".join(PASS_FAIL_OPTIONS)}"',
        allow_blank=True,
    )
    dv_pass_fail.error = "Select PASS, FAIL, UNTESTED, or N/A"
    dv_pass_fail.errorTitle = "Invalid Value"
    dv_pass_fail.prompt = "Select outcome"
    dv_pass_fail.promptTitle = "PASS / FAIL"
    ws.add_data_validation(dv_pass_fail)
    dv_pass_fail.add(f"K2:K{last_row + 10}")

    dv_category = DataValidation(
        type="list",
        formula1=f'"{",".join(CATEGORY_OPTIONS)}"',
        allow_blank=True,
    )
    dv_category.error = "Select a failure category"
    dv_category.errorTitle = "Invalid Category"
    dv_category.prompt = "Select failure category"
    dv_category.promptTitle = "Failure Category"
    ws.add_data_validation(dv_category)
    dv_category.add(f"L2:L{last_row + 10}")

    # Conditional formatting for PASS/FAIL column
    from openpyxl.formatting.rule import CellIsRule
    ws.conditional_formatting.add(
        f"K2:K{last_row + 10}",
        CellIsRule(
            operator="equal",
            formula=['"PASS"'],
            fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
            font=Font(name="Calibri", size=10, bold=True, color="006100"),
        ),
    )
    ws.conditional_formatting.add(
        f"K2:K{last_row + 10}",
        CellIsRule(
            operator="equal",
            formula=['"FAIL"'],
            fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
            font=Font(name="Calibri", size=10, bold=True, color="9C0006"),
        ),
    )
    ws.conditional_formatting.add(
        f"K2:K{last_row + 10}",
        CellIsRule(
            operator="equal",
            formula=['"UNTESTED"'],
            fill=PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            font=Font(name="Calibri", size=10, bold=True, color="9C6500"),
        ),
    )
    ws.conditional_formatting.add(
        f"K2:K{last_row + 10}",
        CellIsRule(
            operator="equal",
            formula=['"N/A"'],
            fill=PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
            font=Font(name="Calibri", size=10, color="595959"),
        ),
    )

    # ── Sheet 2: Summary ─────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")

    summary_headers = ["Metric", "Count", "Percentage"]
    for col_idx, header in enumerate(summary_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 14

    metrics = [
        ("Total Scenarios", f"=COUNTA(Review!C2:C{last_row + 10})", "=B2/B2"),
        ("UNTESTED", f'=COUNTIF(Review!K2:K{last_row + 10},"UNTESTED")', "=B3/B2"),
        ("PASS", f'=COUNTIF(Review!K2:K{last_row + 10},"PASS")', "=B4/B2"),
        ("FAIL", f'=COUNTIF(Review!K2:K{last_row + 10},"FAIL")', "=B5/B2"),
        ("N/A", f'=COUNTIF(Review!K2:K{last_row + 10},"N/A")', "=B6/B2"),
        ("", "", ""),
        ("By Failure Category", "", ""),
        ("Engine Bug", f'=COUNTIF(Review!L2:L{last_row + 10},"Engine Bug")', ""),
        ("Card Data Wrong", f'=COUNTIF(Review!L2:L{last_row + 10},"Card Data Wrong")', ""),
        ("Scenario Broken", f'=COUNTIF(Review!L2:L{last_row + 10},"Scenario Broken")', ""),
        ("Card Not Implemented", f'=COUNTIF(Review!L2:L{last_row + 10},"Card Not Implemented")', ""),
        ("UI/Interface Issue", f'=COUNTIF(Review!L2:L{last_row + 10},"UI/Interface Issue")', ""),
        ("Other", f'=COUNTIF(Review!L2:L{last_row + 10},"Other")', ""),
    ]

    for row_idx, (label, formula, pct) in enumerate(metrics, 2):
        ws2.cell(row=row_idx, column=1, value=label).font = ROW_FONT
        ws2.cell(row=row_idx, column=1).border = THIN_BORDER
        cell_b = ws2.cell(row=row_idx, column=2)
        cell_b.border = THIN_BORDER
        cell_b.alignment = Alignment(horizontal="center")
        if formula:
            cell_b.value = formula
            cell_b.font = Font(name="Calibri", size=10, bold=True)
        cell_c = ws2.cell(row=row_idx, column=3)
        cell_c.border = THIN_BORDER
        if pct:
            cell_c.value = pct
            cell_c.number_format = "0.0%"
            cell_c.font = ROW_FONT

    # Bold section headers
    for r in (2, 8):
        ws2.cell(row=r, column=1).font = Font(name="Calibri", size=10, bold=True)

    ws2.freeze_panes = "A2"

    # ── Sheet 3: Instructions ────────────────────────────────────────
    ws3 = wb.create_sheet("Instructions")

    instructions = [
        ["FIELD", "HOW TO FILL", "EXAMPLE"],
        [],
        [
            "PASS / FAIL",
            "Dropdown. Set to PASS when the card scenario loads and the card behaves correctly per rulebook. Set to FAIL when something is wrong. Use UNTESTED until you've checked it. Use N/A if the scenario is unreachable or the card has no effect in this state.",
            "FAIL",
        ],
        [],
        [
            "Failure Category",
            "Dropdown. Required when FAIL is set. Categorize the root cause:\n"
            "  - Engine Bug: The engine mis-handles a valid card interaction.\n"
            "  - Card Data Wrong: The JSON card definition (execution_model, rules_text, etc.) is incorrect.\n"
            "  - Scenario Broken: The scenario state is unreachable or the card is unplayable.\n"
            "  - Card Not Implemented: The engine has no handler for this card's op/effect.\n"
            "  - UI/Interface Issue: The engine is correct but the CLI display is wrong.\n"
            "  - Other: Something else.",
            "Engine Bug",
        ],
        [],
        [
            "Failure Description",
            "Free text. Describe what goes wrong when you try to play the card or verify its effect. Be specific:\n"
            "  - What action did you attempt?\n"
            "  - What did the engine do?\n"
            "  - What did you expect instead?\n"
            'Example: "Playing the card produces no legal moves, but the rules text says it should allow a supplant."',
            "Card resolves with no legal moves. Should allow supplant of a white troop.",
        ],
        [],
        [
            "Expected Behavior",
            "Free text. Write the correct behavior the card should exhibit, as described by the rulebook or card text. This is the target for the fix.",
            "Priority: capture a white troop on any site. Presence rules apply. No resource gain.",
        ],
        [],
        [
            "Tester",
            "Free text. Your initials or name. Useful for tracking who verified which cards.",
            "MK",
        ],
        [],
        [
            "Date Tested",
            "Free text. Date when you reviewed the scenario. ISO format recommended (YYYY-MM-DD).",
            "2026-06-03",
        ],
        [],
        [
            "Notes",
            "Free text. Any additional context: references to rulebook page, related bug tickets, design decisions.",
            "See rulebook p.12. Same issue as banshee scenario.",
        ],
        [],
        [],
        ["WORKFLOW"],
        [],
        [
            "1. Open the Review sheet.",
            "",
            "",
        ],
        [
            "2. Pick a scenario row (start with UNTESTED).",
            "",
            "",
        ],
        [
            "3. Load the scenario in the CLI:",
            "",
            "",
        ],
        [
            '   just play -- --scenario data/scenarios/cards/101_banshee.json',
            "",
            "",
        ],
        [
            "4. Verify the card under test behaves correctly.",
            "",
            "",
        ],
        [
            "5. Set PASS / FAIL, fill Failure Category + Description if failing.",
            "",
            "",
        ],
        [
            "6. The Summary sheet updates automatically with formula counts.",
            "",
            "",
        ],
    ]

    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 72
    ws3.column_dimensions["C"].width = 56

    for row_idx, row_data in enumerate(instructions, 1):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=value)
            cell.font = ROW_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_idx == 1:
                cell.font = Font(name="Calibri", size=10, bold=True)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx in (29,):
        cell = ws3.cell(row=row_idx, column=1)
        cell.font = Font(name="Calibri", size=12, bold=True)

    ws3.freeze_panes = "A2"

    # ── Save ─────────────────────────────────────────────────────────
    wb.save(str(OUTPUT))
    print(f"Review workbook written: {OUTPUT}")
    print(f"  Scenarios listed: {len(files)}")
    print("  Sheets: Review, Summary, Instructions")


if __name__ == "__main__":
    main()
