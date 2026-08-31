"""Rank the 125 cards in ``data/cards/`` by a composite Card Complexity Score (CCS).

The score is QA-leaning: engine/runtime complexity (execution structure, operational
difficulty, dynamism) is weighted above rules-text and state-footprint complexity. Each
of the five dimensions produces raw points; the weighted sum is scaled to 0-100 against a
documented reference maximum and banded into tiers. Per-dimension sub-scores are reported
so the ranking also serves design review and test prioritization.

Usage:
    python -m scripts.card_complexity_review [--top N] [--card CARD_ID] [--no-artifacts]
"""

from __future__ import annotations

import argparse
import ast
import csv
import io
import json
import re
import sys
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from game_setup.loaders import load_card_catalog  # noqa: E402
from game_setup.types import (  # noqa: E402
    CardCatalog,
    CardDefinition,
    ModalChoiceExecutionModel,
    RepeatChoiceExecutionModel,
    SequenceExecutionModel,
)

DEFAULT_CARD_PATH = ROOT_DIR / "data" / "cards"
DEFAULT_OUTPUT_DIR = ROOT_DIR / "artifacts" / "card_complexity"
DEFAULT_TESTS_DIR = ROOT_DIR / "tests"

OP_TIERS: dict[str, int] = {
    "gain_resource": 0,
    "draw_cards": 0,
    "grant_vp": 0,
    "deploy_troops": 1,
    "place_spy": 1,
    "return_spy": 1,
    "return_unit": 1,
    "force_discard": 1,
    "assassinate_troop": 2,
    "supplant_troop": 2,
    "move_troop": 2,
    "recruit_card": 2,
    "play_card": 2,
    "promote_card": 3,
    "devour": 3,
    "devour_cost": 3,
    "custom_effect": 3,
}
UNKNOWN_OP_TIER = 2

DIMENSION_WEIGHTS: dict[str, float] = {
    "D1": 0.25,
    "D2": 0.25,
    "D3": 0.20,
    "D4": 0.15,
    "D5": 0.15,
}

QUANTITY_POINTS: dict[str, float] = {
    "fixed": 0.0,
    "unspecified": 1.5,
    "variable_repeat": 2.0,
}
UNKNOWN_QUANTITY_POINTS = 2.0

CONDITION_POINTS: dict[str, float] = {
    "focus_check": 1.0,
    "conditional_gate": 1.5,
    "timing_trigger": 1.5,
}
UNKNOWN_CONDITION_POINTS = 1.5

DYNAMIC_TEXT_TOKENS = (
    "for each",
    "for every",
    "if you have",
    "if another",
    "if there",
    "up to",
    "choose exactly one mode",
    "choose one",
    "either",
)

EXECUTION_KIND_POINTS: dict[str, float] = {
    "sequence": 0.0,
    "modal_choice": 2.0,
    "repeat_choice": 3.0,
}

TIER_BANDS = (
    (0.0, "Low"),
    (20.0, "Moderate"),
    (40.0, "High"),
    (60.0, "Very High"),
    (80.0, "Extreme"),
)

REFERENCE_MAX_WEIGHTED = 8.0

_SENTENCE_SPLIT = re.compile(r"[.!?]")

CSV_HEADERS = [
    "rank",
    "card_id",
    "name",
    "aspect",
    "cost",
    "score",
    "tier",
    "test_count",
    "d1_raw",
    "d2_raw",
    "d3_raw",
    "d4_raw",
    "d5_raw",
    "d1_weighted",
    "d2_weighted",
    "d3_weighted",
    "d4_weighted",
    "d5_weighted",
]


@dataclass(frozen=True)
class DimensionScore:
    raw: float
    weighted: float


@dataclass(frozen=True)
class CardComplexityResult:
    card_id: str
    name: str
    aspect: str
    cost: int
    total: float
    tier: str
    test_count: int
    dimensions: dict[str, DimensionScore]
    factors: tuple[str, ...]
    top_factors: tuple[str, ...]
    unmapped: tuple[str, ...]


def tier_label(score: float) -> str:
    """Return the tier band label for a 0-100 score."""
    label = "Low"
    for minimum, band_label in TIER_BANDS:
        if score >= minimum:
            label = band_label
    return label


def _sentence_count(text: str) -> int:
    parts = [part for part in _SENTENCE_SPLIT.split(text) if part.strip()]
    return len(parts)


def _dynamic_token_occurrences(text: str) -> int:
    lowered = text.lower()
    return sum(lowered.count(token) for token in DYNAMIC_TEXT_TOKENS)


def _has_dynamic_metadata(metadata: dict[str, Any]) -> bool:
    if "count_from" in metadata:
        return True
    if metadata.get("requires_focus"):
        return True
    if metadata.get("targeting") == "anywhere":
        return True
    return bool(metadata.get("ignore_presence_requirement"))


def score_card(card: CardDefinition, test_count: int = 0) -> CardComplexityResult:
    """Score one card across the five complexity dimensions."""
    factors: list[tuple[str, float]] = []
    unmapped: list[str] = []

    model = card.execution_model
    options: list[Any] = []
    if isinstance(model, SequenceExecutionModel):
        kind = "sequence"
    elif isinstance(model, ModalChoiceExecutionModel):
        kind = "modal_choice"
        options = list(model.options)
    elif isinstance(model, RepeatChoiceExecutionModel):
        kind = "repeat_choice"
        options = list(model.options)
    else:
        kind = "none"
        factors.append(("no_execution_model", 0.0))

    kind_points = EXECUTION_KIND_POINTS.get(kind, 0.0)
    if kind_points:
        factors.append((f"execution={kind}", kind_points))

    extra_options = float(max(0, len(options) - 1))
    if extra_options:
        factors.append(("options_beyond_first", extra_options))

    flattened = card.actions
    extra_actions = 0.5 * max(0, len(flattened) - 1)
    if extra_actions:
        factors.append(("actions_beyond_first", extra_actions))

    d1_raw = kind_points + extra_options + extra_actions

    op_points = 0.0
    targeting_points = 0.0
    filter_points = 0.0
    optional_points = 0.0
    for action in flattened:
        tier = OP_TIERS.get(action.op)
        if tier is None:
            tier = UNKNOWN_OP_TIER
            unmapped.append(f"op:{action.op}")
        op_points += tier
        if action.target_scope != "self":
            targeting_points += 1.0
        filter_points += len(action.filters)
        if action.optional:
            optional_points += 1.0

    if op_points:
        factors.append(("op_tier_total", op_points))
    if targeting_points:
        factors.append(("targeting", targeting_points))
    if filter_points:
        factors.append(("filters", filter_points))
    if optional_points:
        factors.append(("optional", optional_points))

    d2_raw = op_points + targeting_points + filter_points + optional_points

    quantity_points = 0.0
    timing_points = 0.0
    dynamic_metadata_points = 0.0
    for action in flattened:
        quantity_kind = action.quantity.kind
        quantity = QUANTITY_POINTS.get(quantity_kind)
        if quantity is None:
            quantity = UNKNOWN_QUANTITY_POINTS
            unmapped.append(f"quantity:{quantity_kind}")
        quantity_points += quantity
        if action.timing != "immediate":
            timing_points += 1.0
        if _has_dynamic_metadata(action.metadata):
            dynamic_metadata_points += 1.0

    if quantity_points:
        factors.append(("quantity_dynamism", quantity_points))
    if timing_points:
        factors.append(("timing", timing_points))
    if dynamic_metadata_points:
        factors.append(("dynamic_metadata", dynamic_metadata_points))

    d3_raw = quantity_points + timing_points + dynamic_metadata_points

    condition_points = 0.0
    for condition in card.global_conditions:
        points = CONDITION_POINTS.get(condition.condition_type)
        if points is None:
            points = UNKNOWN_CONDITION_POINTS
            unmapped.append(f"condition:{condition.condition_type}")
        condition_points += points
        factors.append((f"condition={condition.condition_type}", points))

    d4_raw = condition_points

    rules_text = card.rules_text or ""
    extra_sentences = float(max(0, _sentence_count(rules_text) - 1))
    token_points = min(5.0, 0.5 * _dynamic_token_occurrences(rules_text))
    state_points = 0.25 * len(card.state_contract.reads) + 0.5 * len(card.state_contract.writes)
    if extra_sentences:
        factors.append(("sentences_beyond_first", extra_sentences))
    if token_points:
        factors.append(("dynamic_tokens", token_points))
    if state_points:
        factors.append(("state_contract", state_points))

    d5_raw = extra_sentences + token_points + state_points

    raw_by_dimension = {"D1": d1_raw, "D2": d2_raw, "D3": d3_raw, "D4": d4_raw, "D5": d5_raw}
    dimensions = {
        name: DimensionScore(raw=round(raw, 2), weighted=round(raw * DIMENSION_WEIGHTS[name], 3))
        for name, raw in raw_by_dimension.items()
    }
    weighted_total = sum(dimension.weighted for dimension in dimensions.values())
    total = round(100.0 * min(weighted_total, REFERENCE_MAX_WEIGHTED) / REFERENCE_MAX_WEIGHTED, 1)

    ordered_factors = sorted(factors, key=lambda item: (-item[1], item[0]))
    factor_strings = tuple(
        f"{label} (+{points:.2f})" if points else label for label, points in ordered_factors
    )
    top_factors = tuple(label for label, points in ordered_factors[:2] if points)

    return CardComplexityResult(
        card_id=card.card_id,
        name=card.name,
        aspect=card.aspect,
        cost=card.cost,
        total=total,
        tier=tier_label(total),
        test_count=test_count,
        dimensions=dimensions,
        factors=factor_strings,
        top_factors=top_factors,
        unmapped=tuple(sorted(set(unmapped))),
    )


def _count_test_functions(path: Path) -> int:
    """Count ``def test_*`` functions in a single test file (via AST)."""
    try:
        tree = ast.parse(path.read_text(encoding="utf-8"))
    except (OSError, SyntaxError, UnicodeDecodeError):
        return 0
    return sum(
        1
        for node in ast.walk(tree)
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name.startswith("test_")
    )


def _card_id_from_test_stem(stem: str, card_ids: set[str]) -> str | None:
    """Map a ``test_card_*.py`` stem back to a card_id, tolerating ``_<variant>`` suffixes."""
    prefix = "test_card_"
    if not stem.startswith(prefix):
        return None
    suffix = stem[len(prefix):]
    if suffix in card_ids:
        return suffix
    candidates = [card_id for card_id in card_ids if suffix.startswith(f"{card_id}_")]
    if not candidates:
        return None
    return max(candidates, key=len)


def count_card_tests(tests_dir: Path, card_ids: set[str]) -> dict[str, int]:
    """Count ``test_*`` functions in per-card test files for every known card_id."""
    counts: dict[str, int] = {card_id: 0 for card_id in card_ids}
    tests_dir = Path(tests_dir)
    if not tests_dir.is_dir():
        return counts
    for path in tests_dir.rglob("test_card_*.py"):
        card_id = _card_id_from_test_stem(path.stem, card_ids)
        if card_id is None:
            continue
        counts[card_id] += _count_test_functions(path)
    return counts


def score_catalog(catalog: CardCatalog, test_counts: dict[str, int] | None = None) -> list[CardComplexityResult]:
    """Score every card in a catalog (in catalog order)."""
    counts = test_counts or {}
    return [score_card(card, test_count=counts.get(card.card_id, 0)) for card in catalog.cards]


def rank_results(results: Iterable[CardComplexityResult]) -> list[CardComplexityResult]:
    """Rank results by score descending, breaking ties by card_id ascending."""
    return sorted(results, key=lambda result: (-result.total, result.card_id))


def _collect_unmapped(results: list[CardComplexityResult]) -> list[str]:
    tokens: set[str] = set()
    for result in results:
        tokens.update(result.unmapped)
    return sorted(tokens)


def render_console(results: list[CardComplexityResult], top: int = 20) -> str:
    """Render the ranked console table, tier distribution, and unmapped warnings."""
    ranked = rank_results(results)
    lines: list[str] = [
        f"Card Complexity Review ({len(ranked)} cards)",
        "",
    ]

    tier_counts: dict[str, int] = {}
    for result in ranked:
        tier_counts[result.tier] = tier_counts.get(result.tier, 0) + 1
    distribution = ", ".join(f"{tier}: {tier_counts.get(tier, 0)}" for _, tier in TIER_BANDS)
    lines.append(f"Tier distribution: {distribution}")
    lines.append("")

    header = ("Rank", "card_id", "name", "aspect", "cost", "score", "tier", "tests", "top factors")
    rows = [
        (
            str(rank),
            result.card_id,
            result.name,
            result.aspect,
            str(result.cost),
            f"{result.total:.1f}",
            result.tier,
            str(result.test_count),
            "; ".join(result.top_factors),
        )
        for rank, result in enumerate(ranked[:top], 1)
    ]
    widths = [
        max(len(header[i]), *(len(row[i]) for row in rows))
        for i in range(len(header))
    ]
    lines.append("  ".join(header[i].ljust(widths[i]) for i in range(len(header))))
    lines.append("  ".join("-" * width for width in widths))
    for row in rows:
        lines.append("  ".join(row[i].ljust(widths[i]) for i in range(len(row))))

    unmapped = _collect_unmapped(ranked)
    if unmapped:
        lines.extend(["", "Unmapped tokens (catalog drift):"])
        for token in unmapped:
            lines.append(f"  - {token}")
    else:
        lines.extend(["", "Unmapped tokens: none"])

    return "\n".join(lines)


def render_single_card(result: CardComplexityResult) -> str:
    """Render a detailed single-card view."""
    lines = [
        f"{result.name} (`{result.card_id}`)",
        f"  aspect: {result.aspect} | cost: {result.cost} | tests: {result.test_count}",
        f"  score: {result.total:.1f} | tier: {result.tier}",
        "",
    ]
    for name, dimension in result.dimensions.items():
        lines.append(f"  {name}: raw {dimension.raw:.2f} (weighted {dimension.weighted:.3f})")
    lines.append("")
    if result.factors:
        lines.append("  factors:")
        for factor in result.factors:
            lines.append(f"    - {factor}")
    else:
        lines.append("  factors: none")
    if result.unmapped:
        lines.append("  unmapped:")
        for token in result.unmapped:
            lines.append(f"    - {token}")
    return "\n".join(lines)


def render_csv(results: list[CardComplexityResult]) -> str:
    """Render the full ranked results as CSV."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(CSV_HEADERS)
    for rank, result in enumerate(results, 1):
        writer.writerow(
            [
                rank,
                result.card_id,
                result.name,
                result.aspect,
                result.cost,
                f"{result.total:.1f}",
                result.tier,
                result.test_count,
                *(f"{result.dimensions[name].raw:.2f}" for name in ("D1", "D2", "D3", "D4", "D5")),
                *(f"{result.dimensions[name].weighted:.3f}" for name in ("D1", "D2", "D3", "D4", "D5")),
            ]
        )
    return buffer.getvalue()


def render_json(results: list[CardComplexityResult]) -> str:
    """Render the full structured results as JSON."""
    payload = {
        "card_count": len(results),
        "reference_max_weighted": REFERENCE_MAX_WEIGHTED,
        "dimension_weights": DIMENSION_WEIGHTS,
        "tier_bands": [{"min": minimum, "label": label} for minimum, label in TIER_BANDS],
        "results": [
            {
                "card_id": result.card_id,
                "name": result.name,
                "aspect": result.aspect,
                "cost": result.cost,
                "score": result.total,
                "tier": result.tier,
                "test_count": result.test_count,
                "dimensions": {
                    name: {"raw": result.dimensions[name].raw, "weighted": result.dimensions[name].weighted}
                    for name in ("D1", "D2", "D3", "D4", "D5")
                },
                "factors": list(result.factors),
                "top_factors": list(result.top_factors),
                "unmapped": list(result.unmapped),
            }
            for result in results
        ],
    }
    return json.dumps(payload, indent=2) + "\n"


def render_markdown(results: list[CardComplexityResult]) -> str:
    """Render the full ranked report as Markdown."""
    ranked = rank_results(results)
    lines = [
        "# Card Complexity Review",
        "",
        f"Cards scored: {len(ranked)}",
        "",
        "## Methodology",
        "",
        "The Card Complexity Score (CCS) is a QA-leaning composite weighted toward engine/runtime",
        "complexity, with rules-text complexity secondary. Each dimension yields raw points; the",
        f"weighted sum is scaled to 0-100 against a reference maximum of {REFERENCE_MAX_WEIGHTED:.1f}.",
        "",
        "| Dimension | Weight | What it measures |",
        "|---|---:|---|",
        "| D1 Execution structure | 0.25 | execution-model kind, option count, flattened action count |",
        "| D2 Operational difficulty | 0.25 | op tier, optional actions, targeting scope, filters |",
        "| D3 Dynamism | 0.20 | quantity kind, deferred timing, dynamic metadata |",
        "| D4 Conditionality | 0.15 | global condition types |",
        "| D5 Textual & state footprint | 0.15 | sentence count, dynamic tokens, state reads/writes |",
        "",
        "| Tier | Score range |",
        "|---|---:|",
    ]
    for minimum, label in TIER_BANDS:
        upper = "100" if label == "Extreme" else str(next((m - 0.1 for m, _ in TIER_BANDS if m > minimum), 100))
        lines.append(f"| {label} | {minimum:.0f}-{upper} |")

    lines.extend(
        [
            "",
            "| Op tier | Points | Ops |",
            "|---|---:|---|",
            "| 0 | 0 | gain_resource, draw_cards, grant_vp |",
            "| 1 | 1 | deploy_troops, place_spy, return_spy, return_unit, force_discard |",
            "| 2 | 2 | assassinate_troop, supplant_troop, move_troop, recruit_card, play_card |",
            "| 3 | 3 | promote_card, devour, devour_cost, custom_effect |",
            "",
            "| Quantity kind | Points |",
            "|---|---:|",
            "| fixed | 0 |",
            "| unspecified | 1.5 |",
            "| variable_repeat | 2.0 |",
            "",
            "| Condition type | Points |",
            "|---|---:|",
            "| focus_check | 1.0 |",
            "| conditional_gate | 1.5 |",
            "| timing_trigger | 1.5 |",
            "",
            "## Ranked Results",
            "",
            "| rank | card_id | name | aspect | cost | score | tier | tests |",
            "|---:|---|---|---|---:|---:|---:|---:|",
        ]
    )

    for rank, result in enumerate(ranked, 1):
        lines.append(
            f"| {rank} | {result.card_id} | {result.name} | {result.aspect} | "
            f"{result.cost} | {result.total:.1f} | {result.tier} | {result.test_count} |"
        )

    tier_counts: dict[str, int] = {}
    for result in ranked:
        tier_counts[result.tier] = tier_counts.get(result.tier, 0) + 1

    lines.extend(["", "## Tier Distribution", ""])
    for _, label in TIER_BANDS:
        lines.append(f"- {label}: {tier_counts.get(label, 0)}")

    unmapped = _collect_unmapped(ranked)
    lines.extend(["", "## Unmapped Tokens", ""])
    if unmapped:
        for token in unmapped:
            lines.append(f"- `{token}`")
    else:
        lines.append("None.")

    lines.extend(["", "## Full Results", ""])
    for rank, result in enumerate(ranked, 1):
        lines.append(f"### {rank}. {result.name} (`{result.card_id}`)")
        lines.append("")
        lines.append(f"- aspect: `{result.aspect}` | cost: `{result.cost}` | score: `{result.total:.1f}` | tier: `{result.tier}` | tests: `{result.test_count}`")
        for name in ("D1", "D2", "D3", "D4", "D5"):
            dimension = result.dimensions[name]
            lines.append(f"- {name}: raw `{dimension.raw:.2f}` (weighted `{dimension.weighted:.3f}`)")
        if result.factors:
            factors = ", ".join(f"`{factor}`" for factor in result.factors)
            lines.append(f"- factors: {factors}")
        if result.unmapped:
            tokens = ", ".join(f"`{token}`" for token in result.unmapped)
            lines.append(f"- unmapped: {tokens}")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


_XLSX_COLUMN_WIDTHS: dict[str, int] = {
    "rank": 6,
    "card_id": 26,
    "name": 26,
    "aspect": 12,
    "cost": 6,
    "score": 8,
    "tier": 10,
    "test_count": 11,
    "d1_raw": 8,
    "d2_raw": 8,
    "d3_raw": 8,
    "d4_raw": 8,
    "d5_raw": 8,
    "d1_weighted": 12,
    "d2_weighted": 12,
    "d3_weighted": 12,
    "d4_weighted": 12,
    "d5_weighted": 12,
}


def _result_row(result: CardComplexityResult, rank: int) -> list[Any]:
    """Build a typed data row (numbers stay numeric for the spreadsheet)."""
    return [
        rank,
        result.card_id,
        result.name,
        result.aspect,
        result.cost,
        result.total,
        result.tier,
        result.test_count,
        result.dimensions["D1"].raw,
        result.dimensions["D2"].raw,
        result.dimensions["D3"].raw,
        result.dimensions["D4"].raw,
        result.dimensions["D5"].raw,
        result.dimensions["D1"].weighted,
        result.dimensions["D2"].weighted,
        result.dimensions["D3"].weighted,
        result.dimensions["D4"].weighted,
        result.dimensions["D5"].weighted,
    ]


def write_xlsx(results: list[CardComplexityResult], path: Path) -> None:
    """Write the ranked results to an .xlsx workbook (one sheet, header + data rows)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Card Complexity"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for column_index, header in enumerate(CSV_HEADERS, 1):
        cell = worksheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, result in enumerate(results, 2):
        for column_index, value in enumerate(_result_row(result, row_index - 1), 1):
            worksheet.cell(row=row_index, column=column_index, value=value)

    for column_index, header in enumerate(CSV_HEADERS, 1):
        column_letter = worksheet.cell(row=1, column=column_index).column_letter
        worksheet.column_dimensions[column_letter].width = _XLSX_COLUMN_WIDTHS.get(header, 12)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(str(path))


def write_artifacts(results: list[CardComplexityResult], output_dir: Path) -> list[Path]:
    """Write the Markdown, CSV, JSON, and XLSX reports; return the written paths."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    ranked = rank_results(results)
    paths = [
        output_dir / "complexity_review.md",
        output_dir / "complexity_review.csv",
        output_dir / "complexity_review.json",
        output_dir / "complexity_review.xlsx",
    ]
    paths[0].write_text(render_markdown(results), encoding="utf-8")
    paths[1].write_text(render_csv(ranked), encoding="utf-8")
    paths[2].write_text(render_json(ranked), encoding="utf-8")
    write_xlsx(ranked, paths[3])
    return paths


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Rank card complexity across data/cards/.")
    parser.add_argument("--card-path", default=str(DEFAULT_CARD_PATH), help="Card directory (default: data/cards)")
    parser.add_argument("--tests-dir", default=str(DEFAULT_TESTS_DIR), help="Tests directory (default: tests)")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Artifact output directory")
    parser.add_argument("--top", type=int, default=20, help="Number of rows in the console table (default: 20)")
    parser.add_argument("--card", help="Score a single card in detail by card_id")
    parser.add_argument("--no-artifacts", action="store_true", help="Skip writing report artifacts")
    return parser


def main(argv: list[str] | None = None) -> int:
    """CLI entry point: load the catalog, score, rank, and report."""
    args = _build_parser().parse_args(argv)

    card_path = Path(args.card_path)
    try:
        catalog = load_card_catalog(card_path)
    except Exception as exc:  # noqa: BLE001 - surface a clear CLI error
        print(f"error: failed to load card catalog from {card_path}: {exc}", file=sys.stderr)
        return 1

    card_ids = {card.card_id for card in catalog.cards}
    test_counts = count_card_tests(Path(args.tests_dir), card_ids)
    results = score_catalog(catalog, test_counts=test_counts)

    if args.card:
        match = next((result for result in results if result.card_id == args.card), None)
        if match is None:
            print(f"error: no card with card_id '{args.card}'", file=sys.stderr)
            return 1
        print(render_single_card(match))
        return 0

    ranked = rank_results(results)
    print(render_console(ranked, top=args.top))

    if not args.no_artifacts:
        output_dir = Path(args.output_dir)
        written = write_artifacts(ranked, output_dir)
        print("")
        for path in written:
            print(f"wrote {path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
