"""Generate first-deck family registry, normalized catalog, and engine coverage report.

This script treats docs/source/first_deck_review.md as the source of truth.
"""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from difflib import get_close_matches
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
REVIEW_PATH = ROOT / "docs" / "source" / "first_deck_review.md"
WORKBOOK_CATALOG_PATH = ROOT / "data" / "cards" / "catalog.json"
DECK_ROSTER_PATH = ROOT / "data" / "decks" / "first_deck_rosters.json"
GAP_REPORT_PATH = ROOT / "docs" / "generated" / "engine_action_gap_report.md"

FULL_DECK_SPECS: dict[str, dict[str, Any]] = {
    "Aberrations": {"deck_id": "aberrations", "name": "Aberrations", "kind": "full_deck", "total_cards": 40},
    "Dragon": {"deck_id": "dragon", "name": "Dragon", "kind": "full_deck", "total_cards": 40},
    "Drow": {"deck_id": "drow", "name": "Drow", "kind": "full_deck", "total_cards": 40},
    "Elementals": {"deck_id": "elementals", "name": "Elementals", "kind": "full_deck", "total_cards": 40},
    "Fungus": {"deck_id": "fungus", "name": "Fungus", "kind": "full_deck", "total_cards": 40},
    "Undead": {"deck_id": "undead", "name": "Undead", "kind": "full_deck", "total_cards": 40},
}

SPECIAL_STACK_SPECS: dict[str, dict[str, Any]] = {
    "insane_outcast": {
        "deck_id": "insane_outcast",
        "name": "Insane Outcast",
        "kind": "single_card_stack",
        "total_cards": 30,
    },
    "house_guard": {
        "deck_id": "house_guard",
        "name": "House Guard",
        "kind": "single_card_stack",
        "total_cards": 15,
    },
    "priestess_of_lolth": {
        "deck_id": "priestess_of_lolth",
        "name": "Priestess of Lolth",
        "kind": "single_card_stack",
        "total_cards": 15,
    },
}

STARTER_DECK_SPEC: dict[str, Any] = {
    "deck_id": "starter_deck",
    "name": "Starter Deck",
    "kind": "starter_deck",
    "total_cards": 10,
    "per_player": True,
}

LEGACY_BASE_CARDS: list[dict[str, Any]] = []

FIELD_PATTERN = re.compile(r"^- ([a-z_]+): (.*)$")
CARD_HEADER_PATTERN = re.compile(r"^### card_id: (.+)$")


@dataclass(frozen=True)
class ReviewedCard:
    card_id: str
    workbook_card_id: str
    name: str
    cost: int
    deck_vp: int
    inner_circle_vp: int
    aspect: str
    aspect_2: str | None
    plain_english_effect: str
    provisional_family: str
    notes: str


def _load_workbook_rows() -> dict[str, dict[str, Any]]:
    """Load workbook rows via openpyxl and key by normalized snake_case card id."""

    from openpyxl import load_workbook

    workbook = load_workbook(ROOT / "data" / "cards" / "first_deck_review.xlsx", read_only=True)
    sheet = workbook.active
    rows: dict[str, dict[str, Any]] = {}

    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=3).value
        if not name:
            continue

        card_id = _to_card_id(str(name))
        rows[card_id] = {
            "name": str(name).strip(),
            "cost": _to_int(sheet.cell(row=row, column=4).value),
            "deck_vp": _to_int(sheet.cell(row=row, column=5).value),
            "inner_circle_vp": _to_int(sheet.cell(row=row, column=6).value),
            "deck": str(sheet.cell(row=row, column=7).value or "").strip(),
            "market_count": _to_int(sheet.cell(row=row, column=1).value),
            "starter_count": _to_int(sheet.cell(row=row, column=2).value),
            "aspect": str(sheet.cell(row=row, column=11).value or "none").strip().lower(),
            "aspect_2": _normalized_optional(sheet.cell(row=row, column=12).value),
        }

    return rows


def _to_card_id(name: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", name.lower().strip())
    return re.sub(r"_+", "_", normalized).strip("_")


def _to_int(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip()
    if text in {"", "-", "--", "n/a", "N/A"}:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def _normalized_optional(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text.lower()


def _parse_reviewed_cards(workbook_rows: dict[str, dict[str, Any]]) -> list[ReviewedCard]:
    lines = REVIEW_PATH.read_text(encoding="utf-8").splitlines()

    blocks: list[tuple[str, list[str]]] = []
    current_card_id: str | None = None
    body: list[str] = []

    for line in lines:
        header_match = CARD_HEADER_PATTERN.match(line)
        if header_match:
            if current_card_id is not None:
                blocks.append((current_card_id, body))
            current_card_id = header_match.group(1).strip()
            body = []
            continue
        if current_card_id is not None:
            body.append(line)

    if current_card_id is not None:
        blocks.append((current_card_id, body))

    reviewed_cards: list[ReviewedCard] = []
    workbook_alias_index: dict[str, str] = {}
    for workbook_card_id in workbook_rows:
        workbook_alias_index.setdefault(_canonical_key(workbook_card_id), workbook_card_id)

    for card_id, block_lines in blocks:
        fields: dict[str, str] = {}
        for line in block_lines:
            match = FIELD_PATTERN.match(line)
            if match:
                fields[match.group(1)] = match.group(2)

        status = fields.get("status", "").strip()
        if status != "reviewed":
            continue

        resolved_card_id = card_id
        if resolved_card_id not in workbook_rows:
            alias_match = workbook_alias_index.get(_canonical_key(card_id))
            if alias_match is not None:
                resolved_card_id = alias_match
            else:
                fuzzy_matches = get_close_matches(card_id, list(workbook_rows.keys()), n=1, cutoff=0.88)
                if not fuzzy_matches:
                    raise ValueError(f"Card '{card_id}' is reviewed in markdown but missing in workbook rows")
                resolved_card_id = fuzzy_matches[0]

        workbook_meta = workbook_rows[resolved_card_id]
        reviewed_cards.append(
            ReviewedCard(
                card_id=card_id,
                workbook_card_id=resolved_card_id,
                name=workbook_meta["name"],
                cost=workbook_meta["cost"],
                deck_vp=workbook_meta["deck_vp"],
                inner_circle_vp=workbook_meta["inner_circle_vp"],
                aspect=workbook_meta["aspect"],
                aspect_2=workbook_meta["aspect_2"],
                plain_english_effect=fields.get("plain_english_effect", "").strip(),
                provisional_family=fields.get("provisional_family", "").strip(),
                notes=fields.get("notes", "").strip(),
            )
        )

    if len(reviewed_cards) != 125:
        raise ValueError(f"Expected 125 reviewed cards, found {len(reviewed_cards)}")

    return reviewed_cards


def _canonical_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


ACTION_STATE_READS: dict[str, list[str]] = {
    "gain_resource": ["effect_payload"],
    "deploy_troops": ["player.barracks", "board.troop_slots"],
    "assassinate_troop": ["board.troop_slots", "presence_or_target_rules"],
    "supplant_troop": ["board.troop_slots", "presence_or_target_rules"],
    "place_spy": ["player.spies_available", "board.nodes"],
    "return_spy": ["board.spies"],
    "return_unit": ["board.troop_slots", "board.spies"],
    "draw_cards": ["player.deck", "player.discard_pile"],
    "promote_card": ["player.played_cards", "player.hand", "player.discard_pile"],
    "recruit_card": ["market.row", "card_catalog"],
    "devour_cost": ["source_zone", "devour_rules"],
    "force_discard": ["player.hand", "threshold_or_target_rules"],
    "move_troop": ["board.troop_slots", "board.adjacency"],
    "transfer_trophy": ["player.trophy_hall", "target_player.trophy_hall"],
    "play_card": ["player.inner_circle", "card_catalog"],
    "grant_vp": ["vp_scale_source"],
    "conditional_bonus": ["condition_context"],
    "custom_effect": ["custom_rules"],
}


ACTION_STATE_WRITES: dict[str, list[str]] = {
    "gain_resource": ["resource_pool"],
    "deploy_troops": ["board.troop_slots", "player.barracks"],
    "assassinate_troop": ["board.troop_slots", "player.trophy_hall"],
    "supplant_troop": ["board.troop_slots", "player.trophy_hall"],
    "place_spy": ["board.spies", "player.spies_available"],
    "return_spy": ["board.spies", "player.spies_available"],
    "return_unit": ["board.troop_slots", "board.spies", "player.spies_available"],
    "draw_cards": ["player.hand", "player.deck", "player.discard_pile"],
    "promote_card": ["player.inner_circle", "player.played_cards", "player.hand", "player.discard_pile"],
    "recruit_card": ["player.discard_pile", "market.row", "market.deck"],
    "devour_cost": ["devour_zone_or_equivalent", "source_zone"],
    "force_discard": ["player.hand", "player.discard_pile"],
    "move_troop": ["board.troop_slots"],
    "transfer_trophy": ["player.trophy_hall", "target_player.trophy_hall", "board.troop_slots"],
    "play_card": ["player.played_cards"],
    "grant_vp": ["player.vp_tokens_or_score"],
    "conditional_bonus": ["resource_pool_or_followup_actions"],
    "custom_effect": ["custom_state"],
}


def _quantity_hint(fragment: str) -> dict[str, Any]:
    if "up_to" in fragment:
        return {"kind": "up_to", "value": None}
    if "choose_twice" in fragment or "twice" in fragment:
        return {"kind": "fixed", "value": 2}
    if "repeated" in fragment or "multi" in fragment or "mass" in fragment:
        return {"kind": "variable_repeat", "value": None}
    return {"kind": "unspecified", "value": None}


def _new_action(
    op: str,
    source_fragment: str,
    *,
    target_scope: str = "self",
    timing: str = "immediate",
    optional: bool = False,
    filters: list[str] | None = None,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "action_id": "",
        "op": op,
        "target_scope": target_scope,
        "timing": timing,
        "optional": optional,
        "quantity": _quantity_hint(source_fragment),
        "filters": filters or [],
        "source_fragment": source_fragment,
        "metadata": metadata or {},
    }


def _split_sequence_expression(expression: str) -> list[str]:
    parts: list[str] = []
    for plus_part in expression.strip("_").split("_plus_"):
        parts.extend(part for part in plus_part.split("_then_") if part)
    return [part for part in parts if part]


def _with_action_ids(actions: list[dict[str, Any]], prefix: str) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for index, action in enumerate(actions, start=1):
        action_copy = dict(action)
        action_copy["action_id"] = f"{prefix}_{index}"
        output.append(action_copy)
    return output


def _actions_from_fragment(fragment: str) -> list[dict[str, Any]]:
    normalized = fragment.strip("_")
    if not normalized:
        return []

    if "devour" in normalized and ("_for_" in normalized or "_to_" in normalized):
        delimiter = "_for_" if "_for_" in normalized else "_to_"
        lhs, rhs = normalized.split(delimiter, 1)
        actions = [_devour_action(lhs or normalized)]
        for rhs_fragment in _split_sequence_expression(rhs):
            actions.extend(_actions_from_fragment(rhs_fragment))
        return actions

    if normalized == "gain_power_and_influence":
        gain_power_action = _new_action("gain_resource", normalized, metadata={"resource": "power"})
        gain_power_action["quantity"] = {"kind": "fixed", "value": 2}
        gain_influence_action = _new_action("gain_resource", normalized, metadata={"resource": "influence"})
        gain_influence_action["quantity"] = {"kind": "fixed", "value": 2}
        return [gain_power_action, gain_influence_action]

    if "gain_power" in normalized:
        gain_power_action = _new_action("gain_resource", normalized, metadata={"resource": "power"})
        gain_power_action["quantity"] = {"kind": "fixed", "value": 2}
        return [gain_power_action]
    if "gain_influence" in normalized:
        gain_influence_action = _new_action("gain_resource", normalized, metadata={"resource": "influence"})
        gain_influence_action["quantity"] = {"kind": "fixed", "value": 2}
        return [gain_influence_action]
    if "deploy" in normalized:
        return [_new_action("deploy_troops", normalized, target_scope="board_site")]
    if "assassinate" in normalized:
        filters = ["white_troop_only"] if "white" in normalized else []
        return [_new_action("assassinate_troop", normalized, target_scope="board_site", filters=filters)]
    if "supplant" in normalized:
        filters = ["white_troop_only"] if "white" in normalized else []
        return [_new_action("supplant_troop", normalized, target_scope="board_site", filters=filters)]
    if normalized.startswith("spy") or normalized == "spy":
        return [_new_action("place_spy", normalized, target_scope="board_site")]
    if "return_spy" in normalized:
        return [_new_action("return_spy", normalized, target_scope="board_site")]
    if normalized.startswith("return_enemy_unit") or normalized.startswith("return_other_player_unit"):
        return [_new_action("return_unit", normalized, target_scope="opponent_unit")]
    if normalized.startswith("return_unit"):
        return [_new_action("return_unit", normalized, target_scope="self_or_opponent_unit")]
    if "draw" in normalized:
        return [_new_action("draw_cards", normalized)]
    if "promote" in normalized:
        timing = "end_of_turn" if ("triggered" in normalized or "end_of_turn" in normalized) else "immediate"
        return [_new_action("promote_card", normalized, timing=timing)]
    if "recruit" in normalized:
        return [_new_action("recruit_card", normalized, target_scope="market")]
    if "devour" in normalized:
        return [_devour_action(normalized)]
    if "discard" in normalized:
        return [_new_action("force_discard", normalized, target_scope="opponent")]
    if "move_enemy_troop" in normalized or "move_troop" in normalized:
        return [_new_action("move_troop", normalized, target_scope="board_site")]
    if "take" in normalized and "trophy" in normalized:
        return [_new_action("transfer_trophy", normalized, target_scope="opponent_trophy_hall")]
    if "play" in normalized:
        return [_new_action("play_card", normalized, target_scope="inner_circle_or_market")]
    if "vp" in normalized:
        return [_new_action("grant_vp", normalized, target_scope="self")]
    if "focus" in normalized or "conditional" in normalized:
        return [_new_action("conditional_bonus", normalized)]

    return [_new_action("custom_effect", normalized)]


def _devour_action(fragment: str) -> dict[str, Any]:
    source_zone = "unknown"
    if "hand" in fragment:
        source_zone = "hand"
    elif "inner_circle" in fragment:
        source_zone = "inner_circle"
    elif "market" in fragment:
        source_zone = "market"
    elif "self" in fragment:
        source_zone = "played_self"

    return _new_action(
        "devour_cost",
        fragment,
        target_scope=source_zone,
        optional="optional" in fragment,
        metadata={"source_zone": source_zone, "self_replace": "self_replace" in fragment},
    )


def _derive_global_conditions(family_id: str, effect_text: str) -> list[dict[str, Any]]:
    lowered = effect_text.lower()
    conditions: list[dict[str, Any]] = []
    if "focus" in family_id or "focus" in lowered:
        conditions.append(
            {
                "condition_type": "focus_check",
                "description": "Requires matching-aspect support card in hand or played zone",
            }
        )
    if "conditional" in family_id or " if " in f" {lowered} " or "threshold" in family_id:
        conditions.append(
            {
                "condition_type": "conditional_gate",
                "description": "Effect has conditional branch or threshold requirement",
            }
        )
    if "triggered" in family_id or "end_of_turn" in family_id or "end of turn" in lowered:
        conditions.append(
            {
                "condition_type": "timing_trigger",
                "description": "Effect includes deferred or triggered timing",
            }
        )
    return conditions


def _build_execution_model(family_id: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    if family_id.startswith("modal_"):
        modal_body = family_id[len("modal_") :]
        option_expressions = [segment for segment in modal_body.split("_or_") if segment]

        options: list[dict[str, Any]] = []
        flattened_actions: list[dict[str, Any]] = []
        for option_index, expression in enumerate(option_expressions, start=1):
            option_actions_raw: list[dict[str, Any]] = []
            for fragment in _split_sequence_expression(expression):
                option_actions_raw.extend(_actions_from_fragment(fragment))
            option_actions = _with_action_ids(option_actions_raw, f"option_{option_index}_action")
            options.append({"option_id": f"option_{option_index}", "actions": option_actions})
            flattened_actions.extend(option_actions)

        return {
            "kind": "modal_choice",
            "selection": "exactly_one",
            "options": options,
        }, flattened_actions

    if family_id.startswith("choose_twice_"):
        repeat_body = family_id[len("choose_twice_") :]
        option_expressions = [segment for segment in repeat_body.split("_or_") if segment]

        options: list[dict[str, Any]] = []
        flattened_actions: list[dict[str, Any]] = []
        for option_index, expression in enumerate(option_expressions, start=1):
            option_actions_raw: list[dict[str, Any]] = []
            for fragment in _split_sequence_expression(expression):
                option_actions_raw.extend(_actions_from_fragment(fragment))
            option_actions = _with_action_ids(option_actions_raw, f"option_{option_index}_action")
            options.append({"option_id": f"option_{option_index}", "actions": option_actions})
            flattened_actions.extend(option_actions)

        return {
            "kind": "repeat_choice",
            "repeat_count": 2,
            "allow_repeat": True,
            "options": options,
        }, flattened_actions

    sequence_actions_raw: list[dict[str, Any]] = []
    for fragment in _split_sequence_expression(family_id):
        sequence_actions_raw.extend(_actions_from_fragment(fragment))
    sequence_actions = _with_action_ids(sequence_actions_raw, "action")
    return {
        "kind": "sequence",
        "actions": sequence_actions,
    }, sequence_actions


def _state_contract(actions: list[dict[str, Any]]) -> dict[str, list[str]]:
    reads: list[str] = []
    writes: list[str] = []
    for action in actions:
        op = action["op"]
        for read in ACTION_STATE_READS.get(op, []):
            if read not in reads:
                reads.append(read)
        for write in ACTION_STATE_WRITES.get(op, []):
            if write not in writes:
                writes.append(write)
    return {"reads": reads, "writes": writes}


def _is_modal_effect(effect_text: str) -> bool:
    lowered = effect_text.lower()
    if "choose exactly one mode" in lowered or "choose one mode" in lowered:
        return True
    if "either" in lowered and " or " in lowered:
        return True
    return False


def _build_effect_families(cards: list[ReviewedCard]) -> dict[str, Any]:
    by_family: dict[str, list[ReviewedCard]] = defaultdict(list)
    for card in cards:
        by_family[card.provisional_family].append(card)

    families = []
    for family_id in sorted(by_family):
        family_cards = by_family[family_id]
        representative = family_cards[0]
        execution_model, flattened_actions = _build_execution_model(family_id)

        families.append(
            {
                "family_id": family_id,
                "canonical_description": representative.plain_english_effect,
                "machine": {
                    "schema_version": "2.0",
                    "execution_model": execution_model,
                    "actions": flattened_actions,
                    "global_conditions": _derive_global_conditions(family_id, representative.plain_english_effect),
                    "state_contract": _state_contract(flattened_actions),
                    "parameter_schema": {
                        "numeric_values": "optional",
                        "targets": "optional",
                        "thresholds": "optional",
                        "filters": "optional",
                        "timing": "optional",
                    },
                },
                "cards": [
                    {
                        "card_id": card.card_id,
                        "name": card.name,
                    }
                    for card in sorted(family_cards, key=lambda c: c.card_id)
                ],
            }
        )

    return {
        "registry_id": "first_deck_effect_families",
        "source": "docs/source/first_deck_review.md",
        "card_count": len(cards),
        "family_count": len(families),
        "families": families,
    }


def _build_normalized_catalog(cards: list[ReviewedCard]) -> dict[str, Any]:
    normalized_cards: list[dict[str, Any]] = [card.copy() for card in LEGACY_BASE_CARDS]
    seen_card_ids = {card["card_id"] for card in normalized_cards}

    for card in sorted(cards, key=lambda c: c.card_id):
        if card.card_id in seen_card_ids:
            continue
        execution_model, flattened_actions = _build_execution_model(card.provisional_family)

        normalized_cards.append(
            {
                "card_id": card.card_id,
                "name": card.name,
                "cost": card.cost,
                "aspect": card.aspect,
                "secondary_aspects": [card.aspect_2] if card.aspect_2 else [],
                "deck_vp": card.deck_vp,
                "inner_circle_vp": card.inner_circle_vp,
                "rules_text": card.plain_english_effect,
                "notes": card.notes,
                "execution_model": execution_model,
                "actions": flattened_actions,
                "global_conditions": _derive_global_conditions(card.provisional_family, card.plain_english_effect),
                "state_contract": _state_contract(flattened_actions),
                "effect_key": "generic_card",
                "effect_payload": {},
            }
        )
        seen_card_ids.add(card.card_id)

    _apply_catalog_overrides(normalized_cards)

    return {
        "catalog_id": "base_catalog",
        "cards": normalized_cards,
    }


def _apply_catalog_overrides(cards: list[dict[str, Any]]) -> None:
    cards_by_id = {card["card_id"]: card for card in cards}

    def _mark_anywhere(action: dict[str, Any]) -> None:
        action["metadata"] = {
            **action.get("metadata", {}),
            "targeting": "anywhere",
            "ignore_presence_requirement": True,
        }

    def _set_scaled_vp(action: dict[str, Any], *, count_from: str, per: int) -> None:
        action["metadata"] = {
            **action.get("metadata", {}),
            "count_from": count_from,
            "per": per,
        }

    def _sync_sequence_actions(card: dict[str, Any]) -> None:
        card["actions"] = list(card["execution_model"]["actions"])

    def _sync_modal_actions(card: dict[str, Any]) -> None:
        card["actions"] = [
            action
            for option in card["execution_model"]["options"]
            for action in option["actions"]
        ]

    aboleth = cards_by_id["aboleth"]
    option_1_action = dict(aboleth["execution_model"]["options"][0]["actions"][0])
    option_2_action = dict(aboleth["execution_model"]["options"][1]["actions"][0])
    option_1_action_1 = {
        **option_1_action,
        "action_id": "option_1_action_1",
        "quantity": {"kind": "fixed", "value": 1},
    }
    option_1_action_2 = {
        **option_1_action,
        "action_id": "option_1_action_2",
        "quantity": {"kind": "fixed", "value": 1},
    }
    option_2_action_1 = {
        **option_2_action,
        "action_id": "option_2_action_1",
        "metadata": {**option_2_action.get("metadata", {}), "count_from": "spies_on_board"},
    }
    aboleth["execution_model"]["options"][0]["actions"] = [option_1_action_1, option_1_action_2]
    aboleth["execution_model"]["options"][1]["actions"] = [option_2_action_1]
    aboleth["actions"] = [option_1_action_1, option_1_action_2, option_2_action_1]

    advocate = cards_by_id["advocate"]
    advocate_gain = next(dict(action) for action in advocate["actions"] if action["op"] == "gain_resource")
    advocate_promote = next(dict(action) for action in advocate["actions"] if action["op"] == "promote_card")
    advocate_gain["quantity"] = {"kind": "fixed", "value": 2}
    advocate_promote["metadata"] = {
        **advocate_promote.get("metadata", {}),
        "requires_another_played_card": True,
    }
    advocate["execution_model"] = {
        "kind": "modal_choice",
        "selection": "exactly_one",
        "options": [
            {"option_id": "gain_influence", "actions": [advocate_gain]},
            {"option_id": "promote_end_of_turn", "actions": [advocate_promote]},
        ],
    }
    advocate["actions"] = [advocate_gain, advocate_promote]

    balor = cards_by_id["balor"]
    balor["execution_model"]["actions"][1]["source_fragment"] = "supplant_white_troop_anywhere"
    _mark_anywhere(balor["execution_model"]["actions"][1])
    balor["execution_model"]["actions"][2]["quantity"] = {"kind": "fixed", "value": 1}
    _sync_sequence_actions(balor)

    black_dragon = cards_by_id["black_dragon"]
    black_dragon["execution_model"]["actions"][0]["source_fragment"] = "supplant_white_troop_anywhere"
    _mark_anywhere(black_dragon["execution_model"]["actions"][0])
    _set_scaled_vp(
        black_dragon["execution_model"]["actions"][1],
        count_from="trophy_hall_white",
        per=3,
    )
    _sync_sequence_actions(black_dragon)

    blue_dragon = cards_by_id["blue_dragon"]
    blue_dragon["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 2}
    blue_dragon["execution_model"]["actions"][0]["metadata"] = {
        **blue_dragon["execution_model"]["actions"][0].get("metadata", {}),
        "requires_another_played_card": True,
    }
    blue_dragon["execution_model"]["actions"][1] = {
        "action_id": "action_2",
        "op": "grant_vp",
        "target_scope": "self",
        "timing": "end_of_turn",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "scaled_vp_from_promoted_cards",
        "metadata": {
            "count_from": "inner_circle_cards",
            "per": 3,
        },
    }
    _sync_sequence_actions(blue_dragon)
    blue_dragon["state_contract"] = _state_contract(blue_dragon["actions"])

    death_knight = cards_by_id["death_knight"]
    _set_scaled_vp(
        death_knight["execution_model"]["actions"][1],
        count_from="trophy_hall_non_white",
        per=5,
    )
    _sync_sequence_actions(death_knight)

    derro = cards_by_id["derro"]
    derro["execution_model"]["actions"][0]["source_fragment"] = "supplant_white_troop_anywhere"
    _mark_anywhere(derro["execution_model"]["actions"][0])
    _sync_sequence_actions(derro)

    high_priest_of_myrkul = cards_by_id["high_priest_of_myrkul"]
    high_priest_of_myrkul["execution_model"]["actions"][1]["optional"] = True
    high_priest_of_myrkul["execution_model"]["actions"][1]["metadata"] = {
        **high_priest_of_myrkul["execution_model"]["actions"][1].get("metadata", {}),
        "required_secondary_aspect": "undead",
        "repeat_while_targets": True,
    }
    _sync_sequence_actions(high_priest_of_myrkul)
    high_priest_of_myrkul["state_contract"] = _state_contract(high_priest_of_myrkul["actions"])

    green_dragon = cards_by_id["green_dragon"]
    _set_scaled_vp(
        green_dragon["execution_model"]["options"][1]["actions"][2],
        count_from="owned_control_markers",
        per=1,
    )
    _sync_modal_actions(green_dragon)

    aerisi = cards_by_id["aerisi_kalinoth"]
    for action in aerisi["execution_model"]["actions"]:
        if action["op"] == "gain_resource" and str(action.get("metadata", {}).get("resource", "")).lower() == "power":
            action["quantity"] = {"kind": "fixed", "value": 1}
        if action["op"] == "place_spy":
            action["quantity"] = {"kind": "fixed", "value": 1}
        if action["op"] == "recruit_card":
            action["metadata"] = {**action.get("metadata", {}), "required_aspect": "guile", "max_cost": 4}
    aerisi["actions"] = list(aerisi["execution_model"]["actions"])

    air_elemental = cards_by_id["air_elemental"]
    option_1_actions = [dict(action) for action in air_elemental["execution_model"]["options"][0]["actions"]]
    deploy_base = dict(air_elemental["execution_model"]["options"][1]["actions"][0])
    deploy_base.pop("metadata", None)
    return_spy_action = {
        "action_id": "option_2_action_1",
        "op": "return_spy",
        "target_scope": "board_site",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "return_own_spy",
        "metadata": {"spy_owner": "self"},
    }
    deploy_1 = {**deploy_base, "action_id": "option_2_action_2", "source_fragment": "deploy_troops_step_1", "metadata": {}}
    deploy_2 = {**deploy_base, "action_id": "option_2_action_3", "source_fragment": "deploy_troops_step_2", "metadata": {}}
    deploy_3 = {**deploy_base, "action_id": "option_2_action_4", "source_fragment": "deploy_troops_step_3", "metadata": {}}
    deploy_1["quantity"] = {"kind": "fixed", "value": 1}
    deploy_2["quantity"] = {"kind": "fixed", "value": 1}
    deploy_3["quantity"] = {"kind": "fixed", "value": 1}
    focus_draw = {
        "action_id": "option_2_action_5",
        "op": "draw_cards",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "fixed", "value": 1},
        "filters": [],
        "source_fragment": "focus_draw",
        "metadata": {"requires_focus": True, "focus_aspect": "guile"},
    }
    option_2_actions = [return_spy_action, deploy_1, deploy_2, deploy_3, focus_draw]
    air_elemental["execution_model"]["options"][0]["actions"] = option_1_actions
    air_elemental["execution_model"]["options"][1]["actions"] = option_2_actions
    air_elemental["actions"] = [*option_1_actions, *option_2_actions]

    air_elemental_myrmidon = cards_by_id["air_elemental_myrmidon"]
    for action in air_elemental_myrmidon["execution_model"]["actions"]:
        if action["op"] == "promote_card":
            action["metadata"] = {**action.get("metadata", {}), "required_aspect": "obedience"}
    air_elemental_myrmidon["actions"] = list(air_elemental_myrmidon["execution_model"]["actions"])

    ambassador = cards_by_id["ambassador"]
    for action in ambassador["execution_model"]["actions"]:
        if action["op"] == "promote_card":
            action["metadata"] = {**action.get("metadata", {}), "requires_another_played_card": True}
    ambassador["actions"] = list(ambassador["execution_model"]["actions"])

    banshee = cards_by_id["banshee"]
    for action in banshee["execution_model"]["actions"]:
        if action["op"] == "conditional_bonus":
            action["metadata"] = {
                **action.get("metadata", {}),
                "condition": "selected_node_total_spies_at_least",
                "min_spies": 2,
                "resource": "power",
                "amount": 3,
            }
    banshee["actions"] = list(banshee["execution_model"]["actions"])

    black_earth_cultist = cards_by_id["black_earth_cultist"]
    for action in black_earth_cultist["execution_model"]["actions"]:
        if action["op"] == "conditional_bonus":
            action["metadata"] = {
                **action.get("metadata", {}),
                "condition": "focus_aspect_present",
                "focus_aspect": "ambition",
                "resource": "influence",
                "amount": 2,
            }
    black_earth_cultist["actions"] = list(black_earth_cultist["execution_model"]["actions"])

    black_wyrmling = cards_by_id["black_wyrmling"]
    for action in black_wyrmling["execution_model"]["actions"]:
        if action["op"] == "gain_resource" and str(action.get("metadata", {}).get("resource", "")).lower() == "influence":
            action["quantity"] = {"kind": "fixed", "value": 1}
    black_wyrmling["actions"] = list(black_wyrmling["execution_model"]["actions"])

    blue_wyrmling = cards_by_id["blue_wyrmling"]
    for action in blue_wyrmling["execution_model"]["actions"]:
        if action["op"] == "gain_resource" and str(action.get("metadata", {}).get("resource", "")).lower() == "influence":
            action["quantity"] = {"kind": "fixed", "value": 3}
    blue_wyrmling["actions"] = list(blue_wyrmling["execution_model"]["actions"])

    bounty_hunter = cards_by_id["bounty_hunter"]
    for action in bounty_hunter["execution_model"]["actions"]:
        if action["op"] == "gain_resource" and str(action.get("metadata", {}).get("resource", "")).lower() == "power":
            action["quantity"] = {"kind": "fixed", "value": 3}
    bounty_hunter["actions"] = list(bounty_hunter["execution_model"]["actions"])

    brainwashed_slave = cards_by_id["brainwashed_slave"]
    option_1_actions = [dict(action) for action in brainwashed_slave["execution_model"]["options"][0]["actions"]]
    return_spy_action = {
        "action_id": "option_2_action_1",
        "op": "return_spy",
        "target_scope": "board_site",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "return_spy_for_power_and_influence",
        "metadata": {},
    }
    gain_power_action = {
        "action_id": "option_2_action_2",
        "op": "gain_resource",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "fixed", "value": 2},
        "filters": [],
        "source_fragment": "gain_power",
        "metadata": {"resource": "power"},
    }
    gain_influence_action = {
        "action_id": "option_2_action_3",
        "op": "gain_resource",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "fixed", "value": 2},
        "filters": [],
        "source_fragment": "gain_influence",
        "metadata": {"resource": "influence"},
    }
    option_2_actions = [return_spy_action, gain_power_action, gain_influence_action]
    brainwashed_slave["execution_model"]["options"][0]["actions"] = option_1_actions
    brainwashed_slave["execution_model"]["options"][1]["actions"] = option_2_actions
    brainwashed_slave["actions"] = [*option_1_actions, *option_2_actions]

    carrion_crawler = cards_by_id["carrion_crawler"]
    for action in carrion_crawler["execution_model"]["actions"]:
        if action["op"] == "gain_resource" and str(action.get("metadata", {}).get("resource", "")).lower() == "power":
            action["quantity"] = {"kind": "fixed", "value": 3}
    carrion_crawler["actions"] = list(carrion_crawler["execution_model"]["actions"])

    cranium_rats = cards_by_id["cranium_rats"]
    cranium_rats["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 2}
    _sync_sequence_actions(cranium_rats)

    crushing_wave_cultist = cards_by_id["crushing_wave_cultist"]
    crushing_wave_cultist["execution_model"]["actions"] = [
        {
            "action_id": "action_1",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "assassinate_white_troop",
            "metadata": {},
        },
        {
            "action_id": "action_2",
            "op": "deploy_troops",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "focus_deploy_troops",
            "metadata": {"requires_focus": True, "focus_aspect": "conquest"},
        },
    ]
    crushing_wave_cultist["actions"] = list(crushing_wave_cultist["execution_model"]["actions"])

    deathblade = cards_by_id["deathblade"]
    deathblade["execution_model"]["actions"] = [
        {
            "action_id": "action_1",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_1",
            "metadata": {},
        },
        {
            "action_id": "action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_2",
            "metadata": {},
        },
    ]
    deathblade["actions"] = list(deathblade["execution_model"]["actions"])

    dragon_cultist = cards_by_id["dragon_cultist"]
    option_1 = [dict(action) for action in dragon_cultist["execution_model"]["options"][0]["actions"]]
    option_2 = [
        {
            "action_id": "option_2_action_1",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "gain_influence",
            "metadata": {"resource": "influence"},
        }
    ]
    dragon_cultist["execution_model"]["options"][0]["actions"] = option_1
    dragon_cultist["execution_model"]["options"][1]["actions"] = option_2
    dragon_cultist["actions"] = [*option_1, *option_2]

    death_tyrant = cards_by_id["death_tyrant"]
    death_tyrant["execution_model"]["actions"] = [
        {
            "action_id": "action_1",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_1",
            "metadata": {},
        },
        {
            "action_id": "action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_2",
            "metadata": {},
        },
        {
            "action_id": "action_3",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_3",
            "metadata": {},
        },
        {
            "action_id": "action_4",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "gain_influence_per_troop_removed_by_effect",
            "metadata": {"resource": "influence", "count_from": "assassinations_by_source_effect"},
        },
    ]
    death_tyrant["actions"] = list(death_tyrant["execution_model"]["actions"])

    dragonclaw = cards_by_id["dragonclaw"]
    for action in dragonclaw["execution_model"]["actions"]:
        if action["op"] == "conditional_bonus":
            action["metadata"] = {
                "condition": "player_trophy_hall_non_white_at_least",
                "min_trophies": 5,
                "resource": "power",
                "amount": 2,
            }
    dragonclaw["actions"] = list(dragonclaw["execution_model"]["actions"])

    drow_negotiator = cards_by_id["drow_negotiator"]
    drow_negotiator["execution_model"]["actions"][0] = {
        "action_id": "action_1",
        "op": "conditional_bonus",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "conditional_influence_from_promoted_cards",
        "metadata": {
            "condition": "player_inner_circle_at_least",
            "min_cards": 4,
            "resource": "influence",
            "amount": 3,
        },
    }
    drow_negotiator["execution_model"]["actions"][1]["metadata"] = {
        **drow_negotiator["execution_model"]["actions"][1].get("metadata", {}),
        "requires_another_played_card": True,
    }
    drow_negotiator["actions"] = list(drow_negotiator["execution_model"]["actions"])

    enchanter_of_thay = cards_by_id["enchanter_of_thay"]
    enchanter_of_thay["execution_model"] = {
        "kind": "modal_choice",
        "options": [
            {
                "option_id": "option_1",
                "label": "place_spy",
                "actions": [
                    {
                        "action_id": "action_1",
                        "op": "place_spy",
                        "target_scope": "board_site",
                        "timing": "immediate",
                        "optional": False,
                        "quantity": {"kind": "unspecified", "value": None},
                        "filters": [],
                        "source_fragment": "spy",
                        "metadata": {},
                    }
                ],
            },
            {
                "option_id": "option_2",
                "label": "return_spy_gain_power",
                "actions": [
                    {
                        "action_id": "action_1",
                        "op": "return_spy",
                        "target_scope": "board_site",
                        "timing": "immediate",
                        "optional": False,
                        "quantity": {"kind": "unspecified", "value": None},
                        "filters": [],
                        "source_fragment": "return_spy_for_power",
                        "metadata": {"spy_owner": "self"},
                    },
                    {
                        "action_id": "action_2",
                        "op": "gain_resource",
                        "target_scope": "self",
                        "timing": "immediate",
                        "optional": False,
                        "quantity": {"kind": "fixed", "value": 4},
                        "filters": [],
                        "source_fragment": "gain_4_power",
                        "metadata": {"resource": "power"},
                    },
                ],
            },
        ],
    }
    enchanter_of_thay["actions"] = [
        {
            "action_id": "action_1",
            "op": "place_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "spy",
            "metadata": {},
        },
        {
            "action_id": "action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_power",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 4},
            "filters": [],
            "source_fragment": "gain_4_power",
            "metadata": {"resource": "power"},
        },
    ]

    earth_elemental = cards_by_id["earth_elemental"]
    earth_elemental["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 1}
    earth_elemental["execution_model"]["actions"][2]["quantity"] = {"kind": "fixed", "value": 1}
    earth_elemental["execution_model"]["actions"][2]["metadata"] = {
        "requires_focus": True,
        "focus_aspect": "ambition",
    }
    earth_elemental["actions"] = list(earth_elemental["execution_model"]["actions"])

    eternal_flame_cultist = cards_by_id["eternal_flame_cultist"]
    eternal_flame_cultist["execution_model"]["actions"][1]["metadata"] = {
        "condition": "focus_aspect_present",
        "focus_aspect": "malice",
        "resource": "power",
        "amount": 2,
    }
    eternal_flame_cultist["actions"] = list(eternal_flame_cultist["execution_model"]["actions"])

    ettin = cards_by_id["ettin"]
    ettin["execution_model"]["options"][0]["actions"] = [
        {
            "action_id": "option_1_action_1",
            "op": "deploy_troops",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 3},
            "filters": [],
            "source_fragment": "deploy_troops",
            "metadata": {},
        }
    ]
    ettin["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "repeated_white_troop_assassinate",
            "metadata": {},
        },
        {
            "action_id": "option_2_action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "repeated_white_troop_assassinate",
            "metadata": {},
        },
    ]
    ettin["actions"] = [
        *ettin["execution_model"]["options"][0]["actions"],
        *ettin["execution_model"]["options"][1]["actions"],
    ]

    fire_elemental = cards_by_id["fire_elemental"]
    fire_elemental["execution_model"]["options"][0]["actions"] = [
        {
            "action_id": "option_1_action_1",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "gain_power",
            "metadata": {"resource": "power"},
        },
        {
            "action_id": "option_1_action_2",
            "op": "draw_cards",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "focus_draw",
            "metadata": {"requires_focus": True, "focus_aspect": "malice"},
        },
    ]
    fire_elemental["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "gain_influence",
            "metadata": {"resource": "influence"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "draw_cards",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "focus_draw",
            "metadata": {"requires_focus": True, "focus_aspect": "malice"},
        },
    ]
    fire_elemental["actions"] = [
        *fire_elemental["execution_model"]["options"][0]["actions"],
        *fire_elemental["execution_model"]["options"][1]["actions"],
    ]

    gar_shatterkeel = cards_by_id["gar_shatterkeel"]
    gar_shatterkeel["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 3}
    gar_shatterkeel["execution_model"]["actions"][1]["metadata"] = {
        "required_aspect": "conquest",
        "max_cost": 4,
    }
    gar_shatterkeel["actions"] = list(gar_shatterkeel["execution_model"]["actions"])

    gibbering_mouther = cards_by_id["gibbering_mouther"]
    gibbering_mouther["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 2}
    gibbering_mouther["execution_model"]["actions"][1] = {
        "action_id": "action_2",
        "op": "custom_effect",
        "target_scope": "opponent_player",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "give_insane_outcast_to_player_with_presence_at_deployed_site",
        "metadata": {"effect_kind": "give_insane_outcast_to_player_with_presence_on_last_selected_node"},
    }
    gibbering_mouther["actions"] = list(gibbering_mouther["execution_model"]["actions"])

    ghoul = cards_by_id["ghoul"]
    ghoul["execution_model"]["actions"][1]["target_scope"] = "opponent_player"
    ghoul["execution_model"]["actions"][1]["metadata"] = {
        "effect_kind": "give_insane_outcast_to_each_opponent",
    }
    ghoul["actions"] = list(ghoul["execution_model"]["actions"])

    elder_brain = cards_by_id["elder_brain"]
    elder_brain["execution_model"]["actions"][0]["source_fragment"] = "promote_top_of_deck"
    elder_brain["actions"] = list(elder_brain["execution_model"]["actions"])

    glabrezu = cards_by_id["glabrezu"]
    glabrezu["execution_model"]["actions"] = [
        glabrezu["execution_model"]["actions"][0],
        {
            "action_id": "action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_1",
            "metadata": {},
        },
        {
            "action_id": "action_3",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_2",
            "metadata": {},
        },
    ]
    glabrezu["actions"] = list(glabrezu["execution_model"]["actions"])

    grazzt = cards_by_id["grazzt"]
    grazzt["execution_model"]["options"][0]["actions"] = [
        {
            "action_id": "option_1_action_1",
            "op": "place_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "place_spy_step_1",
            "metadata": {},
        },
        {
            "action_id": "option_1_action_2",
            "op": "place_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "place_spy_step_2",
            "metadata": {},
        },
    ]
    grazzt["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_site_supplant",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "supplant_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": ["allow_white_troop"],
            "source_fragment": "multi_supplant",
            "metadata": {"requires_returned_spy_site": True},
        },
    ]
    grazzt["actions"] = [
        *grazzt["execution_model"]["options"][0]["actions"],
        *grazzt["execution_model"]["options"][1]["actions"],
    ]

    green_wyrmling = cards_by_id["green_wyrmling"]
    green_wyrmling["execution_model"]["actions"][1]["metadata"] = {
        "condition": "selected_node_has_other_player_troop",
        "resource": "influence",
        "amount": 2,
    }
    green_wyrmling["actions"] = list(green_wyrmling["execution_model"]["actions"])

    grimlock = cards_by_id["grimlock"]
    grimlock["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 1}
    grimlock["execution_model"]["actions"][1]["quantity"] = {"kind": "fixed", "value": 2}
    _sync_sequence_actions(grimlock)

    howling_hatred_cultist = cards_by_id["howling_hatred_cultist"]
    howling_hatred_cultist["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_influence_with_focus_bonus",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 3},
            "filters": [],
            "source_fragment": "gain_influence",
            "metadata": {"resource": "influence"},
        },
        {
            "action_id": "option_2_action_3",
            "op": "conditional_bonus",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "focus_gain_power",
            "metadata": {
                "condition": "focus_aspect_present",
                "focus_aspect": "guile",
                "resource": "power",
                "amount": 1,
            },
        },
    ]
    howling_hatred_cultist["actions"] = [
        *howling_hatred_cultist["execution_model"]["options"][0]["actions"],
        *howling_hatred_cultist["execution_model"]["options"][1]["actions"],
    ]

    imix = cards_by_id["imix"]
    imix["execution_model"]["actions"] = [
        {
            "action_id": "action_1",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 4},
            "filters": [],
            "source_fragment": "gain_power",
            "metadata": {"resource": "power"},
        },
        {
            "action_id": "action_2",
            "op": "conditional_bonus",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "focus_gain_power",
            "metadata": {
                "condition": "focus_aspect_present",
                "focus_aspect": "malice",
                "resource": "power",
                "amount": 2,
            },
        },
    ]
    imix["actions"] = list(imix["execution_model"]["actions"])

    infiltrator = cards_by_id["infiltrator"]
    infiltrator["execution_model"]["actions"][1] = {
        "action_id": "action_2",
        "op": "conditional_bonus",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "conditional_power_if_enemy_troop_present",
        "metadata": {
            "condition": "selected_node_has_other_player_troop",
            "resource": "power",
            "amount": 1,
        },
    }
    infiltrator["actions"] = list(infiltrator["execution_model"]["actions"])

    myconid_adult = cards_by_id["myconid_adult"]
    myconid_adult["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 2}
    myconid_adult["execution_model"]["actions"][1]["target_scope"] = "opponent_player"
    myconid_adult["execution_model"]["actions"][1]["metadata"] = {
        "effect_kind": "give_insane_outcast_to_selected_player",
    }
    myconid_adult["actions"] = list(myconid_adult["execution_model"]["actions"])

    myconid_sovereign = cards_by_id["myconid_sovereign"]
    myconid_sovereign["execution_model"]["actions"][0]["target_scope"] = "opponent_player"
    myconid_sovereign["execution_model"]["actions"][0]["metadata"] = {
        "effect_kind": "give_insane_outcast_to_selected_player",
    }
    myconid_sovereign["actions"] = list(myconid_sovereign["execution_model"]["actions"])

    information_broker = cards_by_id["information_broker"]
    information_broker["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_cards",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "draw_cards",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 3},
            "filters": [],
            "source_fragment": "draw_cards",
            "metadata": {},
        },
    ]
    information_broker["actions"] = [
        *information_broker["execution_model"]["options"][0]["actions"],
        *information_broker["execution_model"]["options"][1]["actions"],
    ]

    intellect_devourer = cards_by_id["intellect_devourer"]
    intellect_devourer["execution_model"]["options"][0]["actions"][0]["quantity"] = {
        "kind": "fixed",
        "value": 3,
    }
    intellect_devourer["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "return_unit",
            "target_scope": "self_unit",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_units",
            "metadata": {},
        },
        {
            "action_id": "option_2_action_2",
            "op": "return_unit",
            "target_scope": "self_unit",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_units",
            "metadata": {},
        },
    ]
    intellect_devourer["actions"] = [
        *intellect_devourer["execution_model"]["options"][0]["actions"],
        *intellect_devourer["execution_model"]["options"][1]["actions"],
    ]

    kobold = cards_by_id["kobold"]
    kobold["execution_model"]["options"][0]["actions"][0]["quantity"] = {"kind": "fixed", "value": 1}
    _sync_modal_actions(kobold)

    jackalwere = cards_by_id["jackalwere"]
    jackalwere["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_power_and_influence",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "gain_power",
            "metadata": {"resource": "power"},
        },
        {
            "action_id": "option_2_action_3",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "gain_influence",
            "metadata": {"resource": "influence"},
        },
    ]
    jackalwere["actions"] = [
        *jackalwere["execution_model"]["options"][0]["actions"],
        *jackalwere["execution_model"]["options"][1]["actions"],
    ]

    marilith = cards_by_id["marilith"]
    marilith["execution_model"]["actions"][1] = {
        "action_id": "action_2",
        "op": "gain_resource",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "fixed", "value": 5},
        "filters": [],
        "source_fragment": "gain_power",
        "metadata": {"resource": "power"},
    }
    marilith["actions"] = list(marilith["execution_model"]["actions"])

    marlos_urnrayle = cards_by_id["marlos_urnrayle"]
    marlos_urnrayle["execution_model"]["actions"][0]["quantity"] = {
        "kind": "fixed",
        "value": 1,
    }
    marlos_urnrayle["execution_model"]["actions"][2]["filters"] = [
        "aspect_ambition_only",
        "max_cost_4",
    ]
    marlos_urnrayle["execution_model"]["actions"][2]["metadata"] = {
        "required_aspect": "ambition",
        "max_cost": 4,
    }
    marlos_urnrayle["actions"] = list(marlos_urnrayle["execution_model"]["actions"])

    master_of_melee_magthere = cards_by_id["master_of_melee_magthere"]
    master_of_melee_magthere["execution_model"]["options"][0]["actions"][0]["quantity"] = {
        "kind": "fixed",
        "value": 4,
    }
    master_of_melee_magthere["execution_model"]["options"][1]["actions"][0]["source_fragment"] = "supplant_white_troop_anywhere"
    _mark_anywhere(master_of_melee_magthere["execution_model"]["options"][1]["actions"][0])
    _sync_modal_actions(master_of_melee_magthere)

    beholder = cards_by_id["beholder"]
    for action in beholder["execution_model"]["actions"]:
        if action["op"] == "custom_effect":
            action["metadata"] = {
                **action.get("metadata", {}),
                "effect_kind": "scaled_resource_from_player_zone",
                "source_zone": "trophy_hall",
                "per": 3,
                "resource": "power",
            }
    beholder["actions"] = list(beholder["execution_model"]["actions"])

    white_wyrmling = cards_by_id["white_wyrmling"]
    for action in white_wyrmling["execution_model"]["actions"]:
        if action["op"] == "deploy_troops":
            action["quantity"] = {"kind": "fixed", "value": 2}
    white_wyrmling["actions"] = list(white_wyrmling["execution_model"]["actions"])

    mercenary_squad = cards_by_id["mercenary_squad"]
    for action in mercenary_squad["execution_model"]["actions"]:
        if action["op"] == "deploy_troops":
            action["quantity"] = {"kind": "fixed", "value": 3}
    mercenary_squad["actions"] = list(mercenary_squad["execution_model"]["actions"])

    masters_of_sorcere = cards_by_id["masters_of_sorcere"]
    option_1_actions = [
        {
            "action_id": "option_1_action_1",
            "op": "place_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "spy",
            "metadata": {},
        },
        {
            "action_id": "option_1_action_2",
            "op": "place_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "spy",
            "metadata": {},
        },
    ]
    option_2_actions = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_power",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 4},
            "filters": [],
            "source_fragment": "gain_power",
            "metadata": {"resource": "power"},
        },
    ]
    masters_of_sorcere["execution_model"] = {
        "kind": "modal_choice",
        "selection": "exactly_one",
        "options": [
            {"option_id": "option_1", "actions": option_1_actions},
            {"option_id": "option_2", "actions": option_2_actions},
        ],
    }
    masters_of_sorcere["actions"] = [*option_1_actions, *option_2_actions]
    masters_of_sorcere["state_contract"] = _state_contract(masters_of_sorcere["actions"])

    mind_flayer = cards_by_id["mind_flayer"]
    mind_flayer_option_1_actions = [
        {
            "action_id": "option_1_action_1",
            "op": "devour_cost",
            "target_scope": "hand",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "hand_devour",
            "metadata": {"source_zone": "hand", "self_replace": False},
        },
        {
            "action_id": "option_1_action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 3},
            "filters": [],
            "source_fragment": "gain_influence",
            "metadata": {"resource": "influence"},
        },
    ]
    mind_flayer_option_2_actions = [
        {
            "action_id": "option_2_action_1",
            "op": "devour_cost",
            "target_scope": "hand",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "hand_devour",
            "metadata": {"source_zone": "hand", "self_replace": False},
        },
        {
            "action_id": "option_2_action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "assassinate",
            "metadata": {},
        },
    ]
    mind_flayer["execution_model"] = {
        "kind": "modal_choice",
        "selection": "exactly_one",
        "options": [
            {"option_id": "option_1", "actions": mind_flayer_option_1_actions},
            {"option_id": "option_2", "actions": mind_flayer_option_2_actions},
        ],
    }
    mind_flayer["actions"] = [*mind_flayer_option_1_actions, *mind_flayer_option_2_actions]
    mind_flayer["state_contract"] = _state_contract(mind_flayer["actions"])

    minotaur_skeleton = cards_by_id["minotaur_skeleton"]
    minotaur_skeleton["execution_model"]["options"][0]["actions"][0]["quantity"] = {
        "kind": "fixed",
        "value": 3,
    }
    minotaur_skeleton["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "self_devour_multi_white_assassinate_single_site",
            "metadata": {},
        },
        {
            "action_id": "option_2_action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "self_devour_multi_white_assassinate_single_site",
            "metadata": {"requires_last_selected_node": True},
        },
        {
            "action_id": "option_2_action_3",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": True,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "self_devour_multi_white_assassinate_single_site",
            "metadata": {"requires_last_selected_node": True},
        },
    ]
    minotaur_skeleton["actions"] = [
        *minotaur_skeleton["execution_model"]["options"][0]["actions"],
        *minotaur_skeleton["execution_model"]["options"][1]["actions"],
    ]

    mummy_lord = cards_by_id["mummy_lord"]
    mummy_lord["execution_model"]["options"][1]["actions"][0]["metadata"] = {
        **mummy_lord["execution_model"]["options"][1]["actions"][0].get("metadata", {}),
        "effect_kind": "steal_white_trophy_to_board",
    }
    _sync_modal_actions(mummy_lord)
    mummy_lord["state_contract"] = _state_contract(mummy_lord["actions"])

    nalfeshnee = cards_by_id["nalfeshnee"]
    nalfeshnee["execution_model"]["actions"][0]["quantity"] = {
        "kind": "fixed",
        "value": 3,
    }
    nalfeshnee["actions"] = list(nalfeshnee["execution_model"]["actions"])

    necromancer = cards_by_id["necromancer"]
    necromancer["execution_model"]["options"][0]["actions"][0]["quantity"] = {
        "kind": "fixed",
        "value": 3,
    }
    necromancer["actions"] = [
        *necromancer["execution_model"]["options"][0]["actions"],
        *necromancer["execution_model"]["options"][1]["actions"],
    ]

    neogi = cards_by_id["neogi"]
    neogi["execution_model"]["actions"][0]["quantity"] = {
        "kind": "fixed",
        "value": 4,
    }
    neogi["execution_model"]["actions"][1]["timing"] = "end_of_turn"
    neogi["actions"] = list(neogi["execution_model"]["actions"])

    ulitharid = cards_by_id["ulitharid"]
    ulitharid["execution_model"]["actions"][0]["metadata"] = {
        "source_zone": "market",
        "max_cost": 4,
    }
    ulitharid["execution_model"]["actions"][1]["metadata"] = {
        "source_zone": "market",
        "self_replace": False,
        "requires_last_selected_market_slot": True,
    }
    ulitharid["actions"] = list(ulitharid["execution_model"]["actions"])

    insane_outcast = cards_by_id["insane_outcast"]
    insane_outcast["execution_model"]["actions"][0]["metadata"] = {
        "effect_kind": "self_purge_to_supply",
    }
    insane_outcast["actions"] = list(insane_outcast["execution_model"]["actions"])

    matron_mother = cards_by_id["matron_mother"]
    matron_mother["execution_model"]["actions"][0]["metadata"] = {
        "effect_kind": "mill_deck_to_discard",
    }
    matron_mother["actions"] = list(matron_mother["execution_model"]["actions"])

    night_hag = cards_by_id["night_hag"]
    night_hag["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_cards",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "draw_cards",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "draw_cards",
            "metadata": {},
        },
    ]
    night_hag["actions"] = [
        *night_hag["execution_model"]["options"][0]["actions"],
        *night_hag["execution_model"]["options"][1]["actions"],
    ]

    noble = cards_by_id["noble"]
    noble["execution_model"]["actions"][0] = {
        "action_id": "action_1",
        "op": "gain_resource",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "fixed", "value": 1},
        "filters": [],
        "source_fragment": "gain_influence",
        "metadata": {"resource": "influence"},
    }
    noble["actions"] = list(noble["execution_model"]["actions"])

    ogre_zombie = cards_by_id["ogre_zombie"]
    _mark_anywhere(ogre_zombie["execution_model"]["actions"][0])
    _sync_sequence_actions(ogre_zombie)

    olhydra = cards_by_id["olhydra"]
    olhydra["execution_model"] = {
        "kind": "sequence",
        "actions": [
            {
                "action_id": "action_1",
                "op": "supplant_troop",
                "target_scope": "board_site",
                "timing": "immediate",
                "optional": False,
                "quantity": {"kind": "unspecified", "value": None},
                "filters": ["white_troop_only"],
                "source_fragment": "supplant_white_troop_anywhere",
                "metadata": {
                    "targeting": "anywhere",
                    "ignore_presence_requirement": True,
                },
            },
            {
                "action_id": "action_2",
                "op": "deploy_troops",
                "target_scope": "board_site",
                "timing": "immediate",
                "optional": False,
                "quantity": {"kind": "fixed", "value": 2},
                "filters": [],
                "source_fragment": "focus_deploy_troops",
                "metadata": {
                    "requires_focus": True,
                    "focus_aspect": "conquest",
                },
            },
        ],
    }
    _sync_sequence_actions(olhydra)
    olhydra["state_contract"] = _state_contract(olhydra["actions"])

    ogremoch = cards_by_id["ogremoch"]
    ogremoch["execution_model"]["actions"] = [
        {
            "action_id": "action_1",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 2},
            "filters": [],
            "source_fragment": "gain_influence",
            "metadata": {"resource": "influence"},
        },
        {
            "action_id": "action_2",
            "op": "promote_card",
            "target_scope": "self",
            "timing": "end_of_turn",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "triggered_promote",
            "metadata": {"requires_another_played_card": True},
        },
        {
            "action_id": "action_3",
            "op": "promote_card",
            "target_scope": "self",
            "timing": "end_of_turn",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "triggered_promote",
            "metadata": {
                "requires_another_played_card": True,
                "requires_focus": True,
                "focus_aspect": "ambition",
            },
        },
    ]
    _sync_sequence_actions(ogremoch)
    ogremoch["state_contract"] = _state_contract(ogremoch["actions"])

    orcus = cards_by_id["orcus"]
    deploy_action = orcus["execution_model"]["actions"][3]
    orcus["execution_model"]["actions"] = [
        orcus["execution_model"]["actions"][0],
        {
            "action_id": "action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 5},
            "filters": [],
            "source_fragment": "gain_power_from_devour",
            "metadata": {"resource": "power"},
        },
        {
            "action_id": "action_3",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_1",
            "metadata": {},
        },
        {
            "action_id": "action_4",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_2",
            "metadata": {},
        },
        {
            **deploy_action,
            "action_id": "action_5",
        },
    ]
    orcus["actions"] = list(orcus["execution_model"]["actions"])

    rath_modar = cards_by_id["rath_modar"]
    rath_modar["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 2}
    rath_modar["actions"] = list(rath_modar["execution_model"]["actions"])

    ravenous_zombies = cards_by_id["ravenous_zombies"]
    ravenous_zombies["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 1}
    ravenous_zombies["actions"] = list(ravenous_zombies["execution_model"]["actions"])

    revenant = cards_by_id["revenant"]
    promote_action = revenant["execution_model"]["actions"][1]
    revenant["execution_model"]["actions"] = [
        {
            "action_id": "action_1",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_1",
            "metadata": {},
        },
        {
            "action_id": "action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": [],
            "source_fragment": "assassinate_step_2",
            "metadata": {},
        },
        {
            **promote_action,
            "action_id": "action_3",
            "metadata": {
                **promote_action.get("metadata", {}),
                "min_trophies": 8,
                "count_from": "trophy_hall_all",
            },
        },
    ]
    revenant["actions"] = list(revenant["execution_model"]["actions"])

    red_dragon = cards_by_id["red_dragon"]
    red_dragon["execution_model"]["actions"][1] = {
        "action_id": "action_2",
        "op": "return_spy",
        "target_scope": "board_site",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "return_enemy_spy",
        "metadata": {"spy_owner": "opponent", "free_enemy_return": True},
    }
    _set_scaled_vp(
        red_dragon["execution_model"]["actions"][2],
        count_from="controlled_sites",
        per=1,
    )
    _sync_sequence_actions(red_dragon)

    severin_silrajin = cards_by_id["severin_silrajin"]
    severin_silrajin["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 5}
    severin_silrajin["actions"] = list(severin_silrajin["execution_model"]["actions"])

    skeletal_horde = cards_by_id["skeletal_horde"]
    skeletal_horde["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 2}
    skeletal_horde["actions"] = list(skeletal_horde["execution_model"]["actions"])

    soldier = cards_by_id["soldier"]
    soldier["execution_model"]["actions"][0] = {
        "action_id": "action_1",
        "op": "gain_resource",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "fixed", "value": 1},
        "filters": [],
        "source_fragment": "gain_power",
        "metadata": {"resource": "power"},
    }
    soldier["actions"] = list(soldier["execution_model"]["actions"])

    umber_hulk = cards_by_id["umber_hulk"]
    umber_hulk["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 3}
    umber_hulk["actions"] = list(umber_hulk["execution_model"]["actions"])

    underdark_ranger = cards_by_id["underdark_ranger"]
    underdark_ranger["execution_model"]["actions"] = [
        {
            "action_id": "action_1",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "assassinate_white_troop_step_1",
            "metadata": {},
        },
        {
            "action_id": "action_2",
            "op": "assassinate_troop",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 1},
            "filters": ["white_troop_only"],
            "source_fragment": "assassinate_white_troop_step_2",
            "metadata": {},
        },
    ]
    underdark_ranger["actions"] = list(underdark_ranger["execution_model"]["actions"])

    yan_c_bin = cards_by_id["yan_c_bin"]
    yan_c_bin["execution_model"]["actions"][2] = {
        "action_id": "action_3",
        "op": "place_spy",
        "target_scope": "board_site",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "unspecified", "value": None},
        "filters": [],
        "source_fragment": "focus_spy",
        "metadata": {
            "requires_focus": True,
            "focus_aspect": "guile",
        },
    }
    yan_c_bin["actions"] = list(yan_c_bin["execution_model"]["actions"])

    vampire = cards_by_id["vampire"]
    _set_scaled_vp(
        vampire["execution_model"]["options"][1]["actions"][1],
        count_from="inner_circle_cards",
        per=3,
    )
    _sync_modal_actions(vampire)

    vampire_spawn = cards_by_id["vampire_spawn"]
    vampire_spawn["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 1}
    vampire_spawn["actions"] = list(vampire_spawn["execution_model"]["actions"])

    vrock = cards_by_id["vrock"]
    vrock_option_1_actions = [
        {
            "action_id": "option_1_action_1",
            "op": "place_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "spy",
            "metadata": {},
        }
    ]
    vrock_option_2_actions = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_power",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 5},
            "filters": [],
            "source_fragment": "gain_power",
            "metadata": {"resource": "power"},
        },
    ]
    vrock["execution_model"] = {
        "kind": "modal_choice",
        "selection": "exactly_one",
        "options": [
            {"option_id": "option_1", "actions": vrock_option_1_actions},
            {"option_id": "option_2", "actions": vrock_option_2_actions},
        ],
    }
    vrock["actions"] = [*vrock_option_1_actions, *vrock_option_2_actions]
    vrock["state_contract"] = _state_contract(vrock["actions"])

    watcher_of_thay = cards_by_id["watcher_of_thay"]
    watcher_of_thay["execution_model"]["options"][1]["actions"] = [
        {
            "action_id": "option_2_action_1",
            "op": "return_spy",
            "target_scope": "board_site",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "unspecified", "value": None},
            "filters": [],
            "source_fragment": "return_spy_for_influence",
            "metadata": {"spy_owner": "self"},
        },
        {
            "action_id": "option_2_action_2",
            "op": "gain_resource",
            "target_scope": "self",
            "timing": "immediate",
            "optional": False,
            "quantity": {"kind": "fixed", "value": 3},
            "filters": [],
            "source_fragment": "gain_influence",
            "metadata": {"resource": "influence"},
        },
    ]
    watcher_of_thay["actions"] = [
        *watcher_of_thay["execution_model"]["options"][0]["actions"],
        *watcher_of_thay["execution_model"]["options"][1]["actions"],
    ]

    weaponmaster = cards_by_id["weaponmaster"]
    weaponmaster["execution_model"]["options"][0]["actions"][0]["quantity"] = {"kind": "fixed", "value": 1}
    _sync_modal_actions(weaponmaster)

    water_elemental = cards_by_id["water_elemental"]
    water_elemental["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 2}
    water_elemental["actions"] = list(water_elemental["execution_model"]["actions"])

    white_dragon = cards_by_id["white_dragon"]
    white_dragon["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 3}
    _set_scaled_vp(
        white_dragon["execution_model"]["actions"][1],
        count_from="controlled_sites",
        per=2,
    )
    _sync_sequence_actions(white_dragon)

    wyrmspeaker = cards_by_id["wyrmspeaker"]
    wyrmspeaker["execution_model"]["actions"][0]["quantity"] = {"kind": "fixed", "value": 1}
    wyrmspeaker["actions"] = list(wyrmspeaker["execution_model"]["actions"])

    zuggtmoy = cards_by_id["zuggtmoy"]
    zuggtmoy["execution_model"]["actions"][1] = {
        "action_id": "action_2",
        "op": "gain_resource",
        "target_scope": "self",
        "timing": "immediate",
        "optional": False,
        "quantity": {"kind": "fixed", "value": 3},
        "filters": [],
        "source_fragment": "gain_influence",
        "metadata": {"resource": "influence"},
    }
    zuggtmoy["actions"] = list(zuggtmoy["execution_model"]["actions"])


def _build_deck_rosters(cards: list[ReviewedCard], workbook_rows: dict[str, dict[str, Any]]) -> dict[str, Any]:
    entries_by_deck: dict[str, list[dict[str, Any]]] = defaultdict(list)
    special_decks: dict[str, dict[str, Any]] = {}
    starter_entries: list[dict[str, Any]] = []
    normalizations: list[dict[str, Any]] = []

    for card in sorted(cards, key=lambda card: card.card_id):
        card_id = card.card_id
        workbook_row = workbook_rows.get(card.workbook_card_id)
        if workbook_row is None:
            raise ValueError(f"Reviewed card '{card_id}' is missing workbook deck metadata")

        market_count = int(workbook_row.get("market_count", 0))
        starter_count = int(workbook_row.get("starter_count", 0))
        raw_deck = str(workbook_row.get("deck", "")).strip()

        if starter_count > 0:
            starter_entries.append({"card_id": card_id, "count": starter_count})

        if market_count <= 0:
            continue

        if raw_deck == "Starter":
            starter_entries.append({"card_id": card_id, "count": market_count})
            continue

        if raw_deck == "Standard":
            spec = SPECIAL_STACK_SPECS.get(card_id)
            if spec is None:
                raise ValueError(f"Standard-row card '{card_id}' is missing a special stack specification")
            special_decks[card_id] = {
                **spec,
                "entries": [{"card_id": card_id, "count": market_count}],
            }
            continue

        deck_name = raw_deck
        if deck_name not in FULL_DECK_SPECS:
            raise ValueError(f"Card '{card_id}' references unknown deck '{deck_name}'")
        entries_by_deck[deck_name].append({"card_id": card_id, "count": market_count})

    decks: list[dict[str, Any]] = []

    for deck_name in sorted(FULL_DECK_SPECS):
        spec = FULL_DECK_SPECS[deck_name]
        entries = sorted(entries_by_deck.get(deck_name, []), key=lambda entry: entry["card_id"])
        actual_total = sum(entry["count"] for entry in entries)
        if actual_total != spec["total_cards"]:
            raise ValueError(
                f"Deck '{deck_name}' expected {spec['total_cards']} cards, found {actual_total}"
            )
        decks.append({**spec, "entries": entries})

    starter_entries = sorted(starter_entries, key=lambda entry: entry["card_id"])
    starter_total = sum(entry["count"] for entry in starter_entries)
    if starter_total != STARTER_DECK_SPEC["total_cards"]:
        raise ValueError(
            f"Starter deck expected {STARTER_DECK_SPEC['total_cards']} cards, found {starter_total}"
        )
    decks.append({**STARTER_DECK_SPEC, "entries": starter_entries})

    for card_id in sorted(SPECIAL_STACK_SPECS):
        deck = special_decks.get(card_id)
        if deck is None:
            raise ValueError(f"Special stack '{card_id}' is missing from workbook rows")
        actual_total = sum(entry["count"] for entry in deck["entries"])
        expected_total = int(deck["total_cards"])
        if actual_total != expected_total:
            raise ValueError(
                f"Special stack '{card_id}' expected {expected_total} cards, found {actual_total}"
            )
        decks.append(deck)

    return {
        "deck_roster_id": "first_deck_rosters",
        "source": {
            "catalog": "data/cards/catalog.json",
            "review_log": "docs/source/first_deck_review.md",
            "review_workbook": "data/cards/first_deck_review.xlsx",
        },
        "normalizations": normalizations,
        "decks": decks,
    }


def _engine_support_status(family_id: str) -> tuple[str, list[str]]:
    """Coarse static mapping for family support in current engine."""

    reasons: list[str] = []
    status = "unsupported"

    if family_id in {"gain_power", "gain_influence", "noop"}:
        return "supported", ["Registered effect handler exists in engine.rules"]

    if "gain_power" in family_id or "gain_influence" in family_id:
        status = "partial"
        reasons.append("Resource gain primitives exist, but composed family effect_key is not registered")

    if "promote" in family_id:
        status = "partial"
        reasons.append("Promotion queues exist, but family-specific promote handlers are not registered")

    if "deploy" in family_id:
        status = "partial"
        reasons.append("DeployMove exists; many families require multi-step or conditional deploy logic")

    if "assassinate" in family_id:
        status = "partial"
        reasons.append("AssassinateMove exists; family-specific filters and free-effect variants missing")

    if "return_spy" in family_id or "spy" in family_id:
        if status == "unsupported":
            status = "partial"
        reasons.append("ReturnSpyMove exists; targeted spy-placement and modal wrappers are missing")

    hard_unsupported_markers = [
        "devour",
        "supplant",
        "discard",
        "move_enemy",
        "negative",
        "modal",
        "scaled_vp",
    ]
    if any(marker in family_id for marker in hard_unsupported_markers):
        status = "unsupported"
        reasons.append("Requires mechanics not modeled by current move set or effect registry")

    if not reasons:
        reasons.append("No direct handler or move composition path exists for this family")

    return status, reasons


def _build_gap_report(cards: list[ReviewedCard]) -> str:
    family_registry = _build_effect_families(cards)

    lines = [
        "# Engine Action Gap Report",
        "",
        "Source review log: docs/source/first_deck_review.md",
        "",
        "## Summary",
        "",
        "This report maps derived effect families to current engine support in engine/state.py, engine/moves.py, and engine/rules.py.",
        "",
    ]

    status_counter: Counter[str] = Counter()
    gap_rows: list[tuple[str, str, str]] = []
    for family in family_registry["families"]:
        family_id = family["family_id"]
        status, reasons = _engine_support_status(family_id)
        status_counter[status] += 1
        gap_rows.append((family_id, status, "; ".join(reasons)))

    lines.extend(
        [
            f"- Supported: {status_counter['supported']}",
            f"- Partial: {status_counter['partial']}",
            f"- Unsupported: {status_counter['unsupported']}",
            "",
            "## Family Coverage",
            "",
            "| family_id | status | notes |",
            "|---|---|---|",
        ]
    )

    for family_id, status, notes in sorted(gap_rows, key=lambda row: (row[1], row[0])):
        lines.append(f"| {family_id} | {status} | {notes} |")

    lines.extend(
        [
            "",
            "## Flagged Missing Primitives",
            "",
            "- devour-pile state representation",
            "- white troop distinction and targeting",
            "- modal choice move and branching",
            "- targeted spy placement move",
            "- supplant move and rules",
            "- move-enemy-troop move and rules",
            "- discard mechanics and cause tracking",
            "- scaling math helpers for VP and resource effects",
            "- general effect composition and sequencing model",
            "",
            "## Phased Remediation",
            "",
            "1. Add missing state primitives (devour zone, white-troop markers, discard-cause metadata).",
            "2. Add move primitives (modal choice, place spy, supplant, move enemy troop, devour cost action).",
            "3. Add rule handlers for conditional/scaled/discard/devour families.",
            "4. Register effect handlers for highest-frequency families first.",
            "5. Add focused tests per new primitive and per family class.",
        ]
    )

    return "\n".join(lines) + "\n"


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    workbook_rows = _load_workbook_rows()
    reviewed_cards = _parse_reviewed_cards(workbook_rows)

    catalog = _build_normalized_catalog(reviewed_cards)
    deck_rosters = _build_deck_rosters(reviewed_cards, workbook_rows)
    report = _build_gap_report(reviewed_cards)

    _write_json(WORKBOOK_CATALOG_PATH, catalog)
    _write_json(DECK_ROSTER_PATH, deck_rosters)
    GAP_REPORT_PATH.write_text(report, encoding="utf-8")

    print(f"Wrote {WORKBOOK_CATALOG_PATH}")
    print(f"Wrote {DECK_ROSTER_PATH}")
    print(f"Wrote {GAP_REPORT_PATH}")
    print(f"Cards: {len(reviewed_cards)}")


if __name__ == "__main__":
    main()
